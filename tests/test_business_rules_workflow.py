"""Synthetic customer variations through rules, real execution and XLSX QA."""
import copy
import json
import time
from dataclasses import replace

import pandas as pd
import pytest
from openpyxl import load_workbook

from excel_data_toolbox.adaptive_report import build_adaptive_analysis_report
from excel_data_toolbox.business_contracts import evaluate_metric, metric_contract
from excel_data_toolbox.business_rules import BusinessRuleError, prepare_business_data, validate_business_rules
from excel_data_toolbox.business_workflow import BusinessProfileStore, review_business_data


def data():
    return pd.DataFrame({"流水编号": ["A", "A", "B", "C", "D", "E"],
        "业务日期": ["2026-08-01", "2026-08-01", "2026-09-01", "2026-08-04", "2026-08-05", "2026-08-06"],
        "对账日期": ["2026-09-01"] * 6, "门店": ["东", "东", "西", "东", "西", "东"],
        "原始状态": ["完成", "完成", "完成", "取消", "待付", "完成"],
        "Q": ["100", "100", "20", "50", "30", "10%"], "Z": [40, 40, 5, 10, 10, 1]})


def rules(name="销售明细"):
    return {"schema_version": 1, "tables": [{"source_name": name, "role": "fact",
        "rename": {}, "date_field": "业务日期", "duplicate_key": ["流水编号"],
        "statuses": {"原始状态": {"mode": "filter", "valid": ["完成"], "invalid": ["取消", "待付"]}},
        "metrics": {"Q": {"aggregation": "sum", "unit": "元"}, "Z": {"aggregation": "sum", "unit": "元"}},
        "invalid_numeric": "exclude", "derived": [{"field": "毛利", "operator": "subtract", "left": "Q", "right": "Z", "unit": "元"}]}],
        "joins": [], "requirements": ["profitability", "trend", "ranking"]}


def test_policy_is_not_a_silent_repair_and_preserves_source_records():
    original = data()
    before = original.copy(deep=True)
    prepared = prepare_business_data([original], ["销售明细"], rules())
    pd.testing.assert_frame_equal(original, before)
    assert len(prepared.frames[0]) == 2
    assert len(prepared.originals) == 6 and len(prepared.excluded) == 4
    assert set(prepared.excluded["排除源行号"]) == {3, 5, 6, 7}
    assert prepared.audit.iloc[0]["源记录数"] == 6
    assert float(prepared.frames[0]["毛利"].sum()) == 75


def test_compiler_uses_custom_confirmed_metrics_and_selected_time():
    result = build_adaptive_analysis_report([data()], source_names=["销售明细"],
        user_request="分析月度收入、毛利与门店表现", business_rules=rules())
    overview = result.outputs["管理层通用总览"].set_index("指标")
    assert overview.loc["核心指标：Q", "结果"] == 120
    assert overview.loc["核心指标：Z", "结果"] == 45
    assert overview.loc["核心指标：毛利", "结果"] == 75
    assert set(result.outputs["时间趋势"]["月份"]) == {"2026-08", "2026-09"}
    assert result.report["business_rules_applied"]
    assert result.report["request_coverage"]
    assert "规则排除明细" in result.outputs and len(result.outputs) == 15
    assert "partial" in set(q["status"] for q in result.report["request_coverage"])


@pytest.mark.parametrize("mutation", [
    lambda r: r.update(code="print('unsafe')"),
    lambda r: r.update(schema_version=True),
    lambda r: r["tables"][0]["metrics"]["Q"].update(aggregation="exec"),
    lambda r: r["tables"][0]["statuses"]["原始状态"].update(invalid=["完成"]),
    lambda r: r["tables"][0]["derived"][0].update(operator="python"),
    lambda r: r.update(requirements=["arbitrary_python"]),
])
def test_schema_rejects_code_unknown_operations_and_false_approval(mutation):
    policy = rules()
    mutation(policy)
    with pytest.raises(BusinessRuleError):
        validate_business_rules(policy)


def test_duplicate_key_conflicts_and_missing_keys_block_instead_of_keep_first():
    for frame in (data().assign(Z=[40, 99, 5, 10, 10, 1]), data().assign(流水编号=[None, "A", "B", "C", "D", "E"])):
        with pytest.raises(BusinessRuleError, match="业务键"):
            prepare_business_data([frame], ["销售明细"], rules())


def test_unseen_and_missing_status_cannot_be_silently_excluded():
    for value in (None, "支付失败但回调完成", "退款申请已到账"):
        frame = data()
        frame.loc[2, "原始状态"] = value
        with pytest.raises(BusinessRuleError, match="状态"):
            prepare_business_data([frame], ["销售明细"], rules())


def test_dimension_join_is_explicit_many_to_one_and_cannot_duplicate_measures():
    source = pd.DataFrame({"商品编码": ["A", "A", "B"], "销售额": [10, 20, 30]})
    master = pd.DataFrame({"商品编码": ["A", "B"], "商品类别": ["食品", "日用"]})
    policy = {"schema_version": 1, "tables": [{"source_name": "销售", "role": "fact"}, {"source_name": "商品", "role": "dimension"}],
              "joins": [{"left": "销售", "right": "商品", "left_key": "商品编码", "right_key": "商品编码", "fields": ["商品类别"]}]}
    prepared = prepare_business_data([source, master], ["销售", "商品"], policy)
    assert prepared.frames[0]["销售额"].sum() == 60 and len(prepared.frames[0]) == 3
    with pytest.raises(BusinessRuleError, match="唯一"):
        prepare_business_data([source, pd.concat([master, master])], ["销售", "商品"], policy)
    with pytest.raises(BusinessRuleError, match="覆盖"):
        prepare_business_data([source, master.iloc[:1]], ["销售", "商品"], policy)
    policy["tables"][1]["role"] = "fact"
    with pytest.raises(BusinessRuleError, match="两个事实"):
        prepare_business_data([source, master], ["销售", "商品"], policy)


def test_rule_does_not_override_currency_conflicts_or_convert_money():
    policy = rules()
    with pytest.raises(BusinessRuleError, match="单位|币种"):
        prepare_business_data([data().assign(币种="USD")], ["销售明细"], policy)
    policy["tables"][0]["metrics"]["Z"]["unit"] = "万元"
    with pytest.raises(BusinessRuleError, match="同一单位"):
        prepare_business_data([data()], ["销售明细"], policy)


def test_profile_schema_drift_customer_boundary_and_no_raw_values(tmp_path):
    store = BusinessProfileStore(tmp_path)
    saved = store.save("client-a", "月报", rules(), [data()], ["销售明细"])
    disk = next(tmp_path.rglob("*.json")).read_text(encoding="utf-8")
    record = json.loads(disk)
    assert set(record) == {"schema_version", "profile_id", "title", "rules", "schemas"}
    assert "2026-08-01" not in disk and '"100"' not in disk
    reused = store.load_for_review("client-a", saved["profile_id"], [data()], ["新上传__销售明细"])
    assert reused["tables"][0]["source_name"] == "新上传__销售明细"
    assert store.list("client-b") == []
    for owner, identifier, frame in (("client-b", saved["profile_id"], data()),
                                     ("client-a", "../outside", data()),
                                     ("client-a", saved["profile_id"], data().rename(columns={"Q": "Q2"}))):
        with pytest.raises(BusinessRuleError):
            store.load_for_review(owner, identifier, [frame], ["销售明细"])


def test_profile_does_not_guess_between_identical_schemas(tmp_path):
    store = BusinessProfileStore(tmp_path)
    policy = rules("a")
    policy["tables"].append({**copy.deepcopy(policy["tables"][0]), "source_name": "b"})
    profile = store.save("client", "two facts", policy, [data(), data()], ["a", "b"])
    with pytest.raises(BusinessRuleError, match="歧义"):
        store.load_for_review("client", profile["profile_id"], [data(), data()], ["c", "d"])


def test_metric_evaluation_obeys_explicit_semantics_not_header_keywords():
    frame = pd.DataFrame({"X": [20, 40], "Y": [100, 300], "ratio_custom": [.2, .1]})
    context = {"metrics": {"ratio_custom": {"aggregation": "weighted_ratio", "unit": "%", "numerator": "X", "denominator": "Y"}}}
    contract = metric_contract(0, "custom", frame, "ratio_custom", context=context)
    assert contract.status == "confirmed" and evaluate_metric(frame, contract)[0] == .15
    missing = metric_contract(0, "custom", frame, "ratio_custom", context={"metrics": {
        "ratio_custom": {"aggregation": "weighted_ratio", "unit": "%", "numerator": "absent", "denominator": "Y"}}})
    assert missing.status == "unavailable"


def setup_server(monkeypatch, tmp_path):
    import excel_data_toolbox.server as server
    monkeypatch.setattr(server, "TASK_REPOSITORY", server.TaskRepository(tmp_path / "tasks"))
    session = server.AppSession()
    monkeypatch.setattr(server, "SESSION", session)
    table = session.add_table("销售明细", data(), source="synthetic.xlsx", original=True)
    return server, session, object.__new__(server.ToolboxHandler), table


def test_real_server_confirmation_execution_export_and_single_use(monkeypatch, tmp_path):
    server, session, handler, table = setup_server(monkeypatch, tmp_path)
    draft = handler._business_review({"prompt": "分析月度毛利、门店排名，保留异常记录", "table_ids": [table]})
    assert not session.ai_plans and draft["partial_delivery"]["download_url"]
    assert draft["partial_delivery"]["validation_status"] == "passed"
    result = handler._business_confirm({"review_token": draft["review_token"], "rules": rules(),
        "confirmed": True, "customer_scope": "synthetic-customer", "profile_title": "规则回归测试"})
    assert result["saved_profile"] and result["plan_token"]
    executed = handler._ai_execute({"plan_token": result["plan_token"], "confirmed": True})
    assert executed["download_url"] and executed["status"] == "needs_review"
    workbook = load_workbook(session.downloads[executed["download_url"].split("/")[-1]])
    try:
        assert "规则排除明细" in workbook.sheetnames and "规则处理审计" in workbook.sheetnames
        assert workbook["规则排除明细"].max_row == 5
        assert len(workbook["自适应图表看板"]._charts) >= 2
    finally:
        workbook.close()
    pd.testing.assert_frame_equal(session.get(table).frame, data())
    with pytest.raises(server.ApiError, match="已使用"):
        handler._business_confirm({"review_token": draft["review_token"], "rules": rules(), "confirmed": True})
    with pytest.raises(server.ApiError, match="已使用"):
        handler._ai_execute({"plan_token": result["plan_token"], "confirmed": True})


def test_model_cannot_self_approve_business_rules(monkeypatch, tmp_path):
    from excel_data_toolbox.nl_agent import build_table_catalog, validate_plan
    server, session, _, table = setup_server(monkeypatch, tmp_path)
    payload = server._adaptive_report_plan_payload([session.get(table)], "分析销售月度毛利与门店表现")
    payload["steps"][0]["params"]["business_rules"] = rules()
    plan = validate_plan(payload, build_table_catalog({table: data()}))
    with pytest.raises(server.ApiError, match="自批准"):
        session.issue_ai_plan(table_ids=[table], table_signatures=(server._ai_table_signature(session.get(table)),), plan=plan, model="untrusted")
    assert not session.ai_plans


def test_stage_delivery_keeps_raw_evidence_and_blocks_uncertain_money(monkeypatch, tmp_path):
    from excel_data_toolbox.source_guard import detect_generated_workbook
    _, session, handler, table = setup_server(monkeypatch, tmp_path)
    source = session.get(table).frame.copy(deep=True)
    result = handler._local_adaptive_analysis_plan([session.get(table)], "分析月度销售、利润及退款情况", model="local")
    delivery = result["partial_delivery"]
    path = session.downloads[delivery["download_url"].split("/")[-1]]
    assert delivery["status"] == "needs_review" and delivery["review_items"] > 0
    assert not result["auto_execute"] and not session.ai_plans
    assert len(session.tables) == 1  # Never feed the output back into input.
    pd.testing.assert_frame_equal(session.get(table).frame, source)
    book = load_workbook(path)
    try:
        assert book.sheetnames[0] == "阶段交付说明"
        assert {"待确认事项", "语义契约", "原始记录索引", "数据源确认"} <= set(book.sheetnames)
        text = " ".join(str(cell.value) for row in book["阶段交付说明"] for cell in row)
        assert "完整经营结论待核验" in text and "未批准" in text
        assert book["原始记录索引"].max_row >= len(source) + 1
        assert book["原始记录索引"].max_column == 2
        assert "generated-workbook" in (book.properties.keywords or "")
    finally:
        book.close()
    assert detect_generated_workbook(path) is not None


def test_stage_export_failure_never_issues_download_or_approval(monkeypatch, tmp_path):
    server, session, handler, table = setup_server(monkeypatch, tmp_path)
    def fail(*args, **kwargs):
        raise ValueError("synthetic writer failure")
    monkeypatch.setattr(server, "export_verified", fail)
    with pytest.raises(server.ApiError, match="未发布下载"):
        handler._business_review({"prompt": "分析销售月度收入和利润", "table_ids": [table]})
    assert not session.downloads and not session.business_reviews and not session.ai_plans


def test_stage_export_rechecks_snapshot_before_publishing(monkeypatch, tmp_path):
    server, session, handler, table = setup_server(monkeypatch, tmp_path)
    writer = server.export_verified
    def change(*args, **kwargs):
        result = writer(*args, **kwargs)
        session.get(table).frame.loc[0, "Q"] = "900"
        return result
    monkeypatch.setattr(server, "export_verified", change)
    with pytest.raises(server.ApiError, match="数据源在识别期间变化"):
        handler._business_review({"prompt": "分析销售月度收入和利润", "table_ids": [table]})
    assert not session.downloads and not session.business_reviews and not session.ai_plans
    assert not list(session.output_dir.glob("*.xlsx"))


@pytest.mark.parametrize("change", ["data", "reset", "expired"])
def test_confirmation_is_bound_to_task_snapshot_and_expiry(monkeypatch, tmp_path, change):
    server, session, handler, table = setup_server(monkeypatch, tmp_path)
    draft = handler._business_review({"prompt": "分析销售月度毛利与门店表现", "table_ids": [table]})
    token = draft["review_token"]
    if change == "data":
        session.get(table).frame.loc[0, "Q"] = "900"
    elif change == "reset":
        session.reset()
    else:
        session.business_reviews[token] = replace(session.business_reviews[token], expires_at=time.monotonic() - 1)
    with pytest.raises(server.ApiError):
        handler._business_confirm({"review_token": token, "rules": rules(), "confirmed": True})
    assert not session.ai_plans
    # A read-only stage result is allowed before confirmation, not a business
    # transformation. Failed confirmation must not create another delivery.
    assert len(session.downloads) == (0 if change == "reset" else 1)


def test_generic_analysis_enters_review_on_ambiguous_business_rules(monkeypatch, tmp_path):
    _, session, handler, table = setup_server(monkeypatch, tmp_path)
    draft = handler._local_adaptive_analysis_plan([session.get(table)], "分析月度收入、门店利润，输出报表", model="local")
    assert draft["status"] == "clarification" and draft["business_review"]
    assert not draft.get("plan_token") and not session.ai_plans
    assert review_business_data([data()], ["销售明细"], "分析月度收入、门店利润")["questions"]


def test_review_rejects_output_self_ingestion(monkeypatch, tmp_path):
    server, session, handler, _ = setup_server(monkeypatch, tmp_path)
    output = session.add_table("旧输出", pd.DataFrame({"销售额": [1]}), source="result", original=False)
    with pytest.raises(server.ApiError, match="历史输出"):
        handler._business_review({"prompt": "分析销售月度毛利与门店表现", "table_ids": [output]})


def test_audit_column_collisions_preserve_user_values():
    frame = data().assign(排除原因="customer text", 规则来源表="customer source")
    prepared = prepare_business_data([frame], ["销售明细"], rules())
    assert set(prepared.excluded["排除原因"]) == {"customer text"}
    assert "_排除原因" in prepared.excluded and "_规则来源表" in prepared.originals
    with pytest.raises(BusinessRuleError, match="字段冲突"):
        prepare_business_data([frame.assign(_规则来源表="other")], ["销售明细"], rules())


def test_http_review_confirm_execute_download_through_task_registry(monkeypatch, tmp_path):
    """Actual localhost HTTP, not a mocked remote AI or a browser claim."""
    import http.client
    import io
    import threading
    from excel_data_toolbox.session_registry import SessionProxy, SessionRegistry
    server, session, _, table = setup_server(monkeypatch, tmp_path)
    registry = SessionRegistry(lambda identifier: server.AppSession(restore_task_id=identifier), session)
    monkeypatch.setattr(server, "SESSION", SessionProxy(registry))
    httpd = server.ThreadingHTTPServer(("127.0.0.1", 0), server.ToolboxHandler)
    thread = threading.Thread(target=httpd.serve_forever, daemon=True)
    thread.start()
    connection = http.client.HTTPConnection("127.0.0.1", httpd.server_address[1], timeout=30)
    def request(path, payload=None):
        connection.request("POST" if payload is not None else "GET", path,
            json.dumps(payload, ensure_ascii=False).encode("utf-8") if payload is not None else None,
            {"Content-Type": "application/json", "X-Task-ID": session.task_id})
        response = connection.getresponse()
        body = response.read()
        assert response.status == 200, body.decode("utf-8", errors="replace")
        return json.loads(body) if response.headers.get("Content-Type", "").startswith("application/json") else body
    try:
        assert request("/health")["status"] == "ok"
        assert b"businessReviewPanel" in request("/")
        assert b"unifiedRenderBusinessReview" in request("/static/app.js")
        review = request("/api/business-review", {"prompt": "分析月度毛利、门店排名，保留异常记录", "table_ids": [table]})
        plan = request("/api/business-confirm", {"review_token": review["review_token"], "rules": rules(), "confirmed": True})
        result = request("/api/ai-execute", {"plan_token": plan["plan_token"], "confirmed": True})
        contents = request(result["download_url"])
        workbook = load_workbook(io.BytesIO(contents))
        try:
            assert "语义契约" in workbook.sheetnames and "原始记录留档" in workbook.sheetnames
            assert len(workbook["自适应图表看板"]._charts) >= 2
        finally:
            workbook.close()
    finally:
        connection.close()
        httpd.shutdown()
        httpd.server_close()
        thread.join(timeout=5)


def test_custom_ratio_requires_consistent_confirmed_component_units():
    frame = pd.DataFrame({"X": [20, 40], "Y": [100, 300], "R": [.2, .1]})
    context = {"metrics": {"R": {"aggregation": "weighted_ratio", "unit": "%", "numerator": "X", "denominator": "Y"},
                           "X": {"aggregation": "sum", "unit": "元"}, "Y": {"aggregation": "sum", "unit": "万元"}}}
    contract = metric_contract(0, "custom", frame, "R", context=context)
    assert contract.status == "unavailable"


def test_explicit_dictionary_cleans_categories_not_financial_numbers():
    frame = data()
    frame.loc[1, "门店"] = "东店旧名"
    policy = rules()
    policy["tables"][0]["value_maps"] = {"门店": {"东店旧名": "东"}}
    prepared = prepare_business_data([frame], ["销售明细"], policy)
    assert len(prepared.frames[0]) == 2 and len(prepared.excluded) == 4
    assert "东店旧名" in set(prepared.originals["门店"])
    policy["tables"][0]["value_maps"] = {"Q": {"10%": "100"}}
    with pytest.raises(BusinessRuleError, match="重写金额"):
        prepare_business_data([frame], ["销售明细"], policy)


def test_confirmed_units_must_not_still_be_unknown():
    policy = rules()
    policy["tables"][0]["metrics"]["Q"]["unit"] = "单位待确认"
    with pytest.raises(BusinessRuleError, match="待确认"):
        validate_business_rules(policy)


def test_custom_amounts_cannot_hide_metric_specific_currency_declarations():
    frame = pd.DataFrame({"X": [20, 40], "X币种": ["USD", "USD"]})
    context = {"metrics": {"X": {"aggregation": "sum", "unit": "元"}}}
    contract = metric_contract(0, "custom", frame, "X", context=context)
    assert contract.status == "unavailable"
    with pytest.raises(BusinessRuleError, match="币种|单位"):
        prepare_business_data([data().assign(Q币种="USD")], ["销售明细"], rules())


def test_unit_prices_are_not_suggested_or_confirmed_as_additive_money():
    frame = pd.DataFrame({"产品": ["A", "B"], "采购单价": [20, 30]})
    review = review_business_data([frame], ["采购"], "分析价格是否异常")
    assert "采购单价" not in review["draft_rules"]["tables"][0]["metrics"]
    contract = metric_contract(0, "采购", frame, "采购单价", context={"metrics": {
        "采购单价": {"aggregation": "sum", "unit": "元"}}})
    assert contract.status == "unavailable"


def test_confirmed_different_named_dimension_keys_enter_fact_graph():
    facts = pd.DataFrame({"商品编码": ["A", "A", "B"], "销售额": [10, 20, 30]})
    master = pd.DataFrame({"SKU": ["A", "B", "C"], "类别": ["food", "daily", "other"]})
    policy = {"schema_version": 1, "tables": [{"source_name": "订单", "role": "fact"}, {"source_name": "商品", "role": "dimension"}],
              "joins": [{"left": "订单", "right": "商品", "left_key": "商品编码", "right_key": "SKU", "fields": ["类别"]}]}
    result = build_adaptive_analysis_report([facts, master], source_names=["订单", "商品"], business_rules=policy)
    edge = result.outputs["表关系建议"].query("状态 == 'confirmed'").iloc[0]
    assert edge["左侧行覆盖率"] == 1 and edge["右侧键覆盖率"] == pytest.approx(2 / 3)
    assert "商品编码 ↔ SKU" == edge["关联字段"]
