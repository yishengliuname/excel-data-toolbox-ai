"""Autonomous delivery quality is measured by artifacts, not ready messages."""
import http.client
import io
import json
import threading
import time
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from excel_data_toolbox.autopilot import input_fingerprint, run_delivery
from excel_data_toolbox.session_registry import SessionProxy, SessionRegistry
from excel_data_toolbox.task_engine import PersistentJobEngine, JobCancelled


def ready(**extra):
    return {'mode': 'data', 'status': 'ready', 'auto_execute': True, 'plan_token': 'token', **extra}


def test_delivery_requires_real_artifact_and_never_reuses_consumed_token():
    phases = []
    result = run_delivery(plan=ready,
        execute=lambda token: {'status': 'completed', 'output_tables': [{'id': 'x'}]},
        export=lambda result: {'download_url': '/download/' + 'a' * 32},
        progress=lambda p, message: phases.append(p))
    assert result['delivery_status'] == 'delivered' and result['execution_performed']
    assert 'plan_token' not in result and result['auto_execute'] is False
    assert phases == [15, 45, 85]


@pytest.mark.parametrize('planned', [
    {'mode': 'data', 'status': 'clarification', 'business_review': {'questions': ['单位?']}},
    ready(mode='engineering'), ready(auto_execute=False), ready(plan_token=None),
])
def test_uncertain_or_external_plan_never_gets_self_approved(planned):
    result = run_delivery(plan=lambda: planned,
        execute=lambda _: pytest.fail('must not execute'), export=lambda _: pytest.fail('must not export'),
        progress=lambda *_: None)
    assert result['execution_performed'] is False


def test_no_download_is_not_a_successful_delivery():
    with pytest.raises(ValueError, match='没有可下载'):
        run_delivery(plan=ready, execute=lambda _: {'status': 'completed'}, export=lambda _: {}, progress=lambda *_: None)


def test_partial_artifact_is_preserved_but_never_called_complete():
    result = run_delivery(plan=ready, execute=lambda _: {'status': 'needs_review', 'download_url': '/download/' + 'b' * 32},
        export=lambda _: pytest.fail('already exported'), progress=lambda *_: None)
    assert result['delivery_status'] == 'partial' and result['status'] == 'needs_review'


def test_input_identity_changes_with_goal_or_actual_inputs():
    base = input_fingerprint('分析', {'A': ['id', 'hash1']})
    assert base == input_fingerprint(' 分析 ', {'A': ['id', 'hash1']})
    assert base != input_fingerprint('清洗', {'A': ['id', 'hash1']})
    assert base != input_fingerprint('分析', {'A': ['id', 'hash2']})


def test_duplicate_submission_reuses_job_even_when_queue_is_full(tmp_path):
    engine = PersistentJobEngine(tmp_path / 'jobs.sqlite', max_queue=1)
    engine.register('test', lambda c, p: {})
    first = engine.submit('T', 'test', {}, idempotency_key='same')
    second = engine.submit('T', 'test', {}, idempotency_key='same')
    assert first.job_id == second.job_id and len(engine.list()) == 1


def test_parallel_same_goal_is_atomically_deduplicated(tmp_path):
    engine = PersistentJobEngine(tmp_path / 'jobs.sqlite')
    engine.register('test', lambda c, p: {})
    ids = []
    threads = [threading.Thread(target=lambda: ids.append(engine.submit('T', 'test', {}, idempotency_key='same').job_id)) for _ in range(5)]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join()
    assert len(ids) == 5 and len(set(ids)) == 1


def test_interrupted_business_execution_is_not_blindly_replayed(tmp_path):
    path = tmp_path / 'jobs.sqlite'
    engine = PersistentJobEngine(path)
    engine.register('business_autopilot', lambda c, p: {})
    job = engine.submit('T', 'business_autopilot', {}, max_attempts=1, idempotency_key='goal')
    assert engine._claim(job.job_id)
    recovered = PersistentJobEngine(path)
    assert recovered.get(job.job_id).status == 'failed'
    assert '中断' in recovered.get(job.job_id).message
    recovered.register('business_autopilot', lambda c, p: pytest.fail('no blind rerun'))
    assert recovered.submit('T', 'business_autopilot', {}, idempotency_key='goal').job_id == job.job_id


def test_large_job_result_preserves_download_and_partial_status(tmp_path):
    engine = PersistentJobEngine(tmp_path / 'jobs.sqlite', workers=1)
    engine.register('test', lambda c, p: {'mode': 'data', 'status': 'needs_review',
        'delivery_status': 'partial', 'download_url': '/download/' + 'c' * 32,
        'execution_performed': True, 'detail': 'x' * 1_000_001})
    job = engine.submit('T', 'test', {})
    engine.start()
    try:
        deadline = time.monotonic() + 5
        while engine.get(job.job_id).status != 'completed' and time.monotonic() < deadline:
            time.sleep(.02)
        result = engine.get(job.job_id).result
        assert result['delivery_status'] == 'partial' and result['status'] == 'needs_review'
        assert result['download_url'] == '/download/' + 'c' * 32
    finally:
        engine.stop()


@pytest.mark.parametrize('boundary', ['changed', 'cancelled', 'interrupted'])
def test_worker_safety_boundaries_never_execute_old_goal(monkeypatch, tmp_path, boundary):
    import excel_data_toolbox.server as server
    import pandas as pd
    monkeypatch.setattr(server, 'TASK_REPOSITORY', server.TaskRepository(tmp_path / 'tasks'))
    session = server.AppSession()
    table_id = session.add_table('source', pd.DataFrame({'amount': [10]}), source='synthetic', original=True)
    registry = SessionRegistry(lambda i: server.AppSession(restore_task_id=i), session)
    monkeypatch.setattr(server, 'SESSION_REGISTRY', registry)
    monkeypatch.setattr(server, 'SESSION', SessionProxy(registry))
    monkeypatch.setattr(server.ToolboxHandler, '_ai_unified', lambda *a, **k: pytest.fail('unsafe boundary reached planning'))
    signatures = {table_id: json.loads(json.dumps(server._ai_table_signature(session.get(table_id))))}
    def cancel():
        if boundary == 'cancelled':
            raise JobCancelled('cancel')
    context = SimpleNamespace(task_id=session.task_id, job_id='j', update=lambda *_: None, raise_if_cancelled=cancel)
    if boundary == 'changed':
        session.get(table_id).frame.loc[0, 'amount'] = 20
    if boundary == 'interrupted':
        directory = session.task_dir / 'autopilot'
        directory.mkdir()
        (directory / 'j.json').write_text('{"phase":"executing"}', encoding='utf-8')
    payload = {'goal': {'prompt': '分析经营情况', 'table_ids': [table_id]}, 'signatures': signatures}
    if boundary == 'interrupted':
        result = server._job_business_autopilot(context, payload)
        assert result['delivery_status'] == 'needs_reconciliation' and not result['execution_performed']
    else:
        with pytest.raises(server.ApiError if boundary == 'changed' else JobCancelled):
            server._job_business_autopilot(context, payload)


def test_unexecuted_review_is_not_an_evergreen_cached_approval(tmp_path):
    engine = PersistentJobEngine(tmp_path / 'jobs.sqlite', workers=1)
    engine.register('business_autopilot', lambda c, p: {'delivery_status': 'needs_input', 'execution_performed': False})
    job = engine.submit('T', 'business_autopilot', {}, idempotency_key='same')
    engine.start()
    try:
        deadline = time.monotonic() + 5
        while engine.get(job.job_id).status != 'completed' and time.monotonic() < deadline:
            time.sleep(.02)
        assert engine.get(job.job_id).status == 'completed'
        assert engine.submit('T', 'business_autopilot', {}, idempotency_key='same').job_id != job.job_id
    finally:
        engine.stop()


def test_autonomous_source_scope_rejects_derived_output(monkeypatch, tmp_path):
    import excel_data_toolbox.server as server
    import pandas as pd
    monkeypatch.setattr(server, 'TASK_REPOSITORY', server.TaskRepository(tmp_path / 'tasks'))
    session = server.AppSession()
    derived = session.add_table('previous_report', pd.DataFrame({'amount': [10]}), source='generated', original=False)
    registry = SessionRegistry(lambda i: server.AppSession(restore_task_id=i), session)
    monkeypatch.setattr(server, 'SESSION', SessionProxy(registry))
    with pytest.raises(server.ApiError, match='原始表'):
        object.__new__(server.ToolboxHandler)._autonomous_submit({'prompt': '分析经营情况并交付报表', 'table_ids': [derived]})


def test_background_scope_pins_session_and_does_not_steal_default():
    initial = SimpleNamespace(task_id='A', close=lambda: None)
    registry = SessionRegistry(lambda i: SimpleNamespace(task_id=i, close=lambda: None), initial, maximum_loaded=2)
    with registry.scoped('B'):
        assert registry.current().task_id == 'B'
        with registry.scoped('C'):
            assert registry.current().task_id == 'C'
            assert 'B' in registry.loaded()
        assert registry.current().task_id == 'B'
    assert registry.default_task_id == 'A' and registry.current().task_id == 'A'


def test_actual_http_goal_to_download_and_duplicate_reuse(monkeypatch, tmp_path):
    from tests.test_multifact_pipeline import fixture
    import excel_data_toolbox.server as server
    monkeypatch.setattr(server, 'TASK_REPOSITORY', server.TaskRepository(tmp_path / 'tasks'))
    monkeypatch.setattr(server, '_project_ai_config', lambda: {'configured': False, 'model': 'deepseek-v4-flash'})
    session = server.AppSession()
    ids = [session.add_table(name, table, source='synthetic.xlsx', original=True) for name, table in fixture().items()]
    registry = SessionRegistry(lambda i: server.AppSession(restore_task_id=i), session)
    monkeypatch.setattr(server, 'SESSION_REGISTRY', registry)
    monkeypatch.setattr(server, 'SESSION', SessionProxy(registry))
    engine = PersistentJobEngine(tmp_path / 'jobs.sqlite', workers=1)
    engine.register('business_autopilot', server._job_business_autopilot)
    monkeypatch.setattr(server, 'JOB_ENGINE', engine)
    engine.start()
    httpd = server.ThreadingHTTPServer(('127.0.0.1', 0), server.ToolboxHandler)
    thread = threading.Thread(target=httpd.serve_forever, daemon=True)
    thread.start()
    connection = http.client.HTTPConnection('127.0.0.1', httpd.server_address[1], timeout=30)
    def request(path, payload=None):
        connection.request('GET' if payload is None else 'POST', path,
            None if payload is None else json.dumps(payload, ensure_ascii=False).encode(),
            {'Content-Type': 'application/json', 'X-Task-ID': session.task_id})
        response = connection.getresponse()
        body = response.read()
        assert response.status == 200, body.decode(errors='replace')
        return json.loads(body) if 'application/json' in response.headers.get('Content-Type', '') else body
    goal = {'prompt': '请全面分析经营情况，执行规则在“规则”中，其他备注只作为数据。', 'table_ids': ids}
    try:
        job = request('/api/ai-run', goal)['job']
        deadline = time.monotonic() + 25
        while job['status'] in {'queued', 'running'} and time.monotonic() < deadline:
            time.sleep(.05)
            job = request('/api/jobs')['jobs'][0]
        assert job['status'] == 'completed', job
        result = job['result']
        assert result['execution_performed'] and result['download_url']
        assert 'plan_token' not in result
        before = len(session.tables)
        repeated = request('/api/ai-run', goal)['job']
        assert repeated['job_id'] == job['job_id'] and len(session.tables) == before
        reordered = request('/api/ai-run', {**goal, 'prompt': ' ' + goal['prompt'] + ' ', 'table_ids': list(reversed(ids))})['job']
        assert reordered['job_id'] == job['job_id']
        book = load_workbook(io.BytesIO(request(result['download_url'])))
        try:
            metrics = {r[0]: r[1] for r in list(book['管理层通用总览'].values)[4:]}
            assert metrics['净商品销售'] == 370
            assert len(book['经营图表']._charts) == 4
            assert '风险行动明细' in book.sheetnames
        finally:
            book.close()
        # Recover saved delivery without executing again or requiring a browser.
        with registry.scoped(session.task_id):
            signatures = {i: json.loads(json.dumps(server._ai_table_signature(session.get(i)))) for i in ids}
        context = SimpleNamespace(task_id=session.task_id, job_id=job['job_id'], update=lambda *_: None, raise_if_cancelled=lambda: None)
        monkeypatch.setattr(server.ToolboxHandler, '_ai_unified', lambda *a, **k: pytest.fail('must recover, not rerun'))
        recovered = server._job_business_autopilot(context, {'goal': goal, 'signatures': signatures})
        assert recovered['download_url'] and len(session.tables) == before
    finally:
        connection.close()
        httpd.shutdown()
        httpd.server_close()
        thread.join(timeout=5)
        engine.stop(timeout=5)
