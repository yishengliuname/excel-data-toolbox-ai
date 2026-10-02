from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from excel_data_toolbox.delivery_qa import verify_delivery
from excel_data_toolbox.io_utils import export_tables_to_path, load_tables_from_files
from excel_data_toolbox.nl_agent import build_table_catalog, execute_plan, validate_plan
from excel_data_toolbox.customer_order_analysis import (
    build_customer_order_analysis,
    verify_customer_order_workbook,
)


CASE001 = Path(__file__).parent / "fixtures" / "excel_case_001_input.xlsx"
CUSTOMER_REQUEST = (
    "整理8月三家便利店销售数据，删除重复数据，统一门店、商品编码、"
    "商品名称和类别。取消、待支付和明显业务无效数据不计入经营指标，"
    "但要单独保留。考虑退款，计算门店销售、毛利、订单、目标完成率，"
    "并分析商品销售和低毛利商品，生成有汇总、明细和图表的Excel。"
)


def _result():
    if not CASE001.is_file():
        pytest.skip("私有客户 Golden 文件未提供；客户 Excel 不提交到公开仓库")
    tables = load_tables_from_files(CASE001)
    return build_customer_order_analysis(
        list(tables.values()),
        source_names=list(tables.keys()),
        user_request=CUSTOMER_REQUEST,
    )


def test_case001_business_results_match_ground_truth() -> None:
    result = _result()
    report = result.report

    assert report["raw_rows"] == 1125
    assert report["duplicate_rows"] == 28
    assert report["deduplicated_rows"] == 1097
    assert report["valid_rows"] == 1030
    assert report["excluded_rows"] == 95
    assert report["net_sales"] == pytest.approx(12131.87, abs=0.01)
    assert report["gross_profit"] == pytest.approx(5534.47, abs=0.01)
    assert report["gross_margin"] == pytest.approx(0.4562, abs=0.0001)
    assert report["valid_orders"] == 681
    assert report["refund_amount"] == pytest.approx(426.02, abs=0.01)
    assert set(report["stores"]) == {"盐城一店", "盐城二店", "亭湖店"}


def test_case001_outputs_are_canonical_and_auditable() -> None:
    result = _result()
    required = {
        "管理看板",
        "清洗明细",
        "门店汇总",
        "商品分析",
        "异常数据",
        "每日趋势",
        "执行审计",
    }
    assert required.issubset(result.outputs)

    cleaned = result.outputs["清洗明细"]
    assert len(cleaned) == 1030
    assert set(cleaned["门店"].dropna().unique()) == {"盐城一店", "盐城二店", "亭湖店"}
    assert cleaned["商品编码"].str.fullmatch(r"P\d{3}").all()
    assert cleaned["商品名称"].notna().all()
    assert cleaned["类目"].notna().all()
    assert cleaned["毛利"].sum() == pytest.approx(5534.47, abs=0.01)

    excluded = result.outputs["异常数据"]
    assert len(excluded) == 95
    assert excluded["排除原因"].notna().all()
    assert excluded["异常类型"].eq("业务排除").all()

    daily = result.outputs["每日趋势"]
    assert len(daily) > 1
    assert daily["日期"].nunique() > 1


def test_case001_exclusion_reasons_match_business_rules() -> None:
    result = _result()
    reasons = result.outputs["异常数据"]["排除原因"].value_counts().to_dict()
    assert reasons == {
        "标准化后完全重复": 28,
        "订单状态=已取消": 29,
        "订单状态=待支付": 23,
        "门店无法识别": 6,
        "商品编码无法匹配主数据": 5,
        "数量<=0": 3,
        "单价<=0": 1,
    }


def test_case001_export_is_reopenable_styled_and_charted(tmp_path: Path) -> None:
    result = _result()
    output = tmp_path / "case001_订单经营分析报告.xlsx"
    export_tables_to_path(result.outputs, output, include_log=False, overwrite=True)

    acceptance = verify_delivery(output, result.outputs, allow_extra_tables=False)
    assert acceptance.status == "passed"
    workbook_acceptance = verify_customer_order_workbook(str(output))
    assert workbook_acceptance["status"] == "passed"
    assert workbook_acceptance["net_sales"] == pytest.approx(12131.87)
    workbook = load_workbook(output, data_only=False)
    assert list(workbook.sheetnames) == list(result.outputs)
    assert len(workbook["图表看板"]._charts) == 4
    assert workbook["管理看板"]["A1"].value == "订单经营管理看板"
    assert "report_kind=customer_order_analysis" in str(workbook.properties.keywords)


def test_generic_order_rules_prefer_negative_status_and_aggregate_refunds() -> None:
    sales = pd.DataFrame(
        [
            ["O1", "2026-08-01", "一店", "p001", "旧名", "旧类别", 2, 10, 0, "已完成"],
            ["O2", "2026-08-01", "一店", "P001", "旧名", "旧类别", 1, 10, 0, "未完成"],
            ["O3", "2026-08-02", "一店", "P001", "旧名", "旧类别", 20, 10, 0, "已完成"],
            ["O4", "2026-08-03", "一店", "P001", "旧名", "旧类别", 1, 10, 0, "已完成"],
            ["O5", "2026-08-04", "一店", "P001", "旧名", "旧类别", 1, 10, 0, "已完成"],
        ],
        columns=["订单号", "下单时间", "门店名称", "SKU", "商品名称", "商品类别", "数量", "单价", "优惠金额", "订单状态"],
    )
    master = pd.DataFrame(
        [["P001", "正式商品", "食品", 4]], columns=["商品编码", "标准商品名称", "标准类目", "成本价"]
    )
    refunds = pd.DataFrame([["O1", 1], ["O1", 2]], columns=["订单编号", "退款金额"])
    targets = pd.DataFrame([["一店", 1000]], columns=["门店", "销售额目标"])

    result = build_customer_order_analysis(
        [sales, master, refunds, targets],
        source_names=["销售", "商品主数据", "退款", "目标"],
        user_request="整理订单、考虑退款并生成经营分析",
    )

    cleaned = result.outputs["清洗明细"]
    assert set(cleaned["订单号"]) == {"O1", "O3", "O4", "O5"}
    assert cleaned.loc[cleaned["订单号"].eq("O1"), "退款金额"].sum() == pytest.approx(3.0)
    assert cleaned["商品名称"].eq("正式商品").all()
    assert result.outputs["异常数据"]["排除原因"].tolist() == ["订单状态=未完成"]
    assert not result.outputs["统计预警"].empty


def test_customer_order_operation_runs_through_validated_agent_plan() -> None:
    if not CASE001.is_file():
        pytest.skip("私有客户 Golden 文件未提供；客户 Excel 不提交到公开仓库")
    tables = load_tables_from_files(CASE001)
    catalog = build_table_catalog(tables)
    payload = {
        "schema_version": 1,
        "status": "ready",
        "summary": "清洗订单并生成门店、商品和退款经营报告",
        "message": "将执行本地确定性订单分析。",
        "clarification_questions": [],
        "assumptions": [],
        "warnings": [],
        "steps": [
            {
                "id": "order_analysis",
                "operation": "customer_order_analysis",
                "input_ids": list(tables),
                "output_name": "订单经营分析报告",
                "params": {"source_names": list(tables), "user_request": CUSTOMER_REQUEST},
            }
        ],
    }

    plan = validate_plan(payload, catalog)
    result = execute_plan(plan, tables, dry_run=False)

    assert plan.status == "ready"
    assert result.reports["order_analysis"]["validation_status"] == "passed"
    assert result.tables["管理看板"].iloc[0]["结果"] == pytest.approx(12131.87)
