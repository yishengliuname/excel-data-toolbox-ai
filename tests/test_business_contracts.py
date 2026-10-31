"""Synthetic adversarial cases: no new customer or industry-specific template."""
import pandas as pd
import pytest

from excel_data_toolbox.adaptive_report import build_adaptive_analysis_report
from excel_data_toolbox.analysis_compiler import compile_analysis
from excel_data_toolbox.business_contracts import (
    metric_contract, evaluate_metric, relationship_graph, candidate_grain,
    numeric_values,
)


def sales():
    return pd.DataFrame({"订单号": ["A", "B", "C", "D"],
        "日期": ["2026-01-01", "2026-01-02", "2026-02-01", "2026-02-02"],
        "地区": ["东", "西", "东", "西"], "销售额": [100, 200, 300, 400]})


@pytest.mark.parametrize("bad", [None, "not a number", float("inf"), float("-inf")])
def test_incomplete_amount_does_not_become_zero_or_partial_total(bad):
    frame = sales()
    frame["销售额"] = frame["销售额"].astype(object)
    frame.loc[1, "销售额"] = bad
    contract = metric_contract(0, "orders", frame, "销售额")
    assert contract.status == "unavailable"
    assert pd.isna(evaluate_metric(frame, contract)[0])
    report = build_adaptive_analysis_report([frame], source_names=["orders"])
    row = report.outputs["管理层通用总览"].set_index("指标").loc["核心指标：销售额"]
    assert row["结果"] == "不可计算"
    assert "指标不可计算" in report.outputs["异常数据"]["异常类型"].tolist()


def test_duplicates_are_preserved_and_not_silently_aggregated():
    raw = pd.concat([sales(), sales().iloc[[0]]], ignore_index=True)
    result = build_adaptive_analysis_report([raw], source_names=["orders"])
    assert len(result.outputs["主数据分析"]) == len(raw)
    assert len(result.outputs["事实域明细"]) == len(raw)
    assert result.outputs["语义契约"].iloc[0]["状态"] == "unavailable"
    assert result.outputs["异常数据"]["异常类型"].eq("疑似重复（保留）").sum() == 2


@pytest.mark.parametrize("header", ["采购单价", "平均销售额", "周转天数", "每人销售额"])
def test_nonadditive_unit_metrics_are_not_summed(header):
    frame = sales().rename(columns={"销售额": header})
    assert metric_contract(0, "fact", frame, header).status == "unavailable"


@pytest.mark.parametrize("unit_column,units", [
    ("币种", ["CNY", "USD", "CNY", "CNY"]),
    ("采购单位", ["箱", "kg", "kg", "kg"]),
    ("数量单位", ["kg", None, "kg", "kg"]),
])
def test_mixed_or_missing_units_require_conversion_rules(unit_column, units):
    frame = sales().assign(**{unit_column: units})
    assert metric_contract(0, "fact", frame, "销售额").status == "unavailable"


def test_multiple_time_axes_never_select_the_first_column_by_accident():
    frame = sales().assign(退款日期=["2026-03-01"] * 4)
    plan = compile_analysis([frame], source_names=["orders"], user_request="分析月度趋势")
    assert not any(a.kind == "trend" for a in plan.analyses)
    result = build_adaptive_analysis_report([frame], source_names=["orders"], user_request="月度趋势")
    assert result.outputs["时间趋势"].empty
    assert result.outputs["问题与证据"].set_index("分析问题").loc["trend", "证据状态"] == "unavailable"


def test_balance_sums_entities_at_one_latest_snapshot_not_one_last_row():
    frame = pd.DataFrame({"商品编码": ["A", "B", "A", "B"],
        "日期": ["2026-01-01", "2026-01-01", "2026-02-01", "2026-02-01"],
        "库存余额": [10, 20, 30, 40]})
    contract = metric_contract(0, "snapshots", frame, "库存余额")
    assert evaluate_metric(frame, contract)[0] == 70
    missing_latest = frame.iloc[:-1]
    assert metric_contract(0, "snapshots", missing_latest, "库存余额").status == "unavailable"


def test_weighted_ratio_needs_the_same_complete_population_and_unit_scale():
    frame = sales().assign(利润=[30, 60, 90, 40], 利润率=[.3, .3, .3, .1])
    contract = metric_contract(0, "fact", frame, "利润率")
    assert evaluate_metric(frame, contract)[0] == pytest.approx(.22)
    frame.loc[1, "利润"] = None
    assert metric_contract(0, "fact", frame, "利润率").status == "unavailable"
    frame["利润"] = [30, 60, 90, 40]
    scaled = frame.rename(columns={"利润": "利润(万元)", "销售额": "销售额(元)"})
    assert metric_contract(0, "fact", scaled, "利润率").status == "unavailable"


def test_known_money_format_and_scale_do_not_corrupt_values():
    frame = sales().rename(columns={"销售额": "销售额(万元)"})
    frame["销售额(万元)"] = ["1,000", "200", "300", "400"]
    contract = metric_contract(0, "fact", frame, "销售额(万元)")
    assert contract.unit == "万元"
    assert evaluate_metric(frame, contract)[0] == 1900


def test_same_column_name_is_not_a_relationship_or_a_profit_formula():
    income = sales()
    expenses = pd.DataFrame({"日期": ["2026-01-01", "2026-02-01"], "费用金额": [20, 30]})
    result = build_adaptive_analysis_report([income, expenses], source_names=["orders", "expenses"], user_request="公司赚不赚钱，关联分析利润")
    assert result.outputs["表关系建议"].empty  # shared date is not a business key
    assert result.outputs["问题与证据"].set_index("分析问题").loc["profitability", "证据状态"] == "unavailable"
    assert not result.outputs["管理层通用总览"]["指标"].str.contains("净利润").any()


def test_many_to_many_keys_and_bidirectional_coverage_are_explicit():
    left = pd.DataFrame({"订单号": ["A", "A", "B", "C", "C"], "销售额": [1] * 5})
    right = pd.DataFrame({"订单号": ["A", "A", "X"], "退款金额": [1] * 3})
    edge = relationship_graph([left, right], ["sales", "refunds"], ["fact", "fact"])[0]
    assert edge.cardinality == "many_to_many" and edge.status == "blocked"
    assert edge.left_row_coverage == .4
    assert edge.right_row_coverage == pytest.approx(2/3)
    assert edge.left_key_coverage == pytest.approx(1/3)


def test_secondary_fact_can_drive_ranking_and_trend_without_primary_metrics():
    unknown = sales().rename(columns={"销售额": "未定义测量"})
    unknown["说明"] = ["alpha"] * 4  # make this the largest display anchor
    secondary = sales().rename(columns={"订单号": "退款单号", "销售额": "退款金额"})
    plan = compile_analysis([unknown, secondary], source_names=["unknown", "refunds"])
    assert plan.primary_index == 0
    assert any(a.kind == "trend" and a.table_index == 1 for a in plan.analyses)
    assert not any(a.kind == "trend" and a.table_index == 0 for a in plan.analyses)
    result = build_adaptive_analysis_report([unknown, secondary], source_names=["unknown", "refunds"])
    assert set(result.outputs["时间趋势"]["来源事实表"]) == {"refunds"}
    assert result.outputs["自适应图表看板"]["趋势来源"].eq("refunds").all()


def test_top_n_share_uses_full_population_and_chart_keeps_other_categories():
    frame = pd.DataFrame({"订单号": list("ABCDE"), "日期": ["2026-01-01"] * 5,
        "产品": list("ABCDE"), "销售额": [50, 40, 30, 20, 10]})
    result = build_adaptive_analysis_report([frame], source_names=["orders"], top_n=3)
    ranking = result.outputs["分类排名"]
    assert ranking["占比"].sum() == pytest.approx(.8)
    dashboard = result.outputs["自适应图表看板"]
    assert dashboard.loc[dashboard["结构分类"].eq("其他类别"), "结构指标值"].iloc[0] == pytest.approx(30)
    assert dashboard["结构指标值"].sum() == pytest.approx(150)


def test_summary_only_input_gets_structure_audit_not_business_kpis():
    frame = pd.DataFrame({"指标": ["销售额", "利润"], "结果": [100, 20]})
    result = build_adaptive_analysis_report([frame], source_names=["历史管理层总览"])
    assert result.report["fact_count"] == 0
    assert result.outputs["语义契约"].empty
    assert not result.outputs["管理层通用总览"]["指标"].str.startswith("核心指标：").any()


def test_physical_key_is_validated_and_text_is_not_truncated():
    frame = sales()
    frame["备注"] = ["长描述" * 200] * 4
    original = frame.copy(deep=True)
    result = build_adaptive_analysis_report([frame], source_names=["orders"])
    assert len(result.outputs["主数据分析"].iloc[0]["备注"]) == 600
    assert candidate_grain(frame) == ("订单号",)
    pd.testing.assert_frame_equal(original, frame)
    duplicate = pd.concat([frame, frame.iloc[[0]]])
    assert candidate_grain(duplicate) == ()


def test_stateful_rows_require_explicit_population_rules():
    frame = sales().assign(状态=["已完成", "未完成", "支付失败", "已完成"])
    contract = metric_contract(0, "orders", frame, "销售额")
    assert contract.status == "unavailable"
    assert "状态" in str(contract.issues)


@pytest.mark.parametrize("values, expected", [
    (["10%", "20%"], [.1, .2]),
    (["10.0%", "20.5%"], [.1, .205]),
    ([" 10％ ", "-5%", "200"], [.1, -.05, 200]),
])
def test_percentage_notation_is_dtype_independent(values, expected):
    original = pd.Series(values)
    assert numeric_values(original).tolist() == pytest.approx(expected)
    assert original.tolist() == values


@pytest.mark.parametrize("value", ["1%2", "10%%", "%10", "10%foo", "￥10%"])
def test_malformed_percentage_is_not_laundered_into_a_number(value):
    assert pd.isna(numeric_values(pd.Series([value])).iloc[0])


@pytest.mark.parametrize("field", ["销售额", "数量", "满意度评分", "库存余额"])
@pytest.mark.parametrize("other", ["200", "200.5"])
def test_percent_in_nonratio_is_blocked_and_preserved_in_report(field, other):
    frame = sales().rename(columns={"销售额": field})
    frame[field] = ["10%", other, "300", "400"]
    original = frame.copy(deep=True)
    contract = metric_contract(0, "orders", frame, field)
    assert contract.status == "unavailable"
    assert "百分号" in str(contract.issues)
    assert pd.isna(evaluate_metric(frame, contract)[0])
    result = build_adaptive_analysis_report([frame], source_names=["orders"])
    row = result.outputs["管理层通用总览"].set_index("指标").loc[f"核心指标：{field}"]
    assert row["结果"] == "不可计算"
    assert result.outputs["主数据分析"].iloc[0][field] == "10%"
    assert result.outputs["时间趋势"].empty
    pd.testing.assert_frame_equal(original, frame)


def test_integer_percentage_ratio_reaches_compiler_and_report():
    frame = sales().assign(利润=[10, 40, 90, 160], 利润率=["10%", "20%", "30%", "40%"])
    result = build_adaptive_analysis_report([frame], source_names=["orders"])
    assert result.outputs["管理层通用总览"].set_index("指标").loc["核心指标：利润率", "结果"] == pytest.approx(.3)
    assert result.outputs["主数据分析"]["利润率"].tolist() == pytest.approx([.1, .2, .3, .4])


def test_percent_unit_in_amount_header_is_not_a_monetary_amount():
    frame = sales().rename(columns={"销售额": "销售额(%)"})
    assert metric_contract(0, "orders", frame, "销售额(%)").status == "unavailable"


@pytest.mark.parametrize("column,currency,unit", [
    ("币种", "USD", "USD"), ("currency", "eur", "EUR"),
    ("币种", "美元", "USD"), ("销售额币种", "HKD", "HKD"),
    ("币种", "人民币", "元"),
])
def test_declared_single_currency_is_preserved_in_all_report_layers(column, currency, unit):
    frame = sales().assign(**{column: currency})
    original = frame.copy(deep=True)
    contract = metric_contract(0, "orders", frame, "销售额")
    assert contract.status == "inferred" and contract.unit == unit
    assert evaluate_metric(frame, contract)[0] == 1000
    plan = compile_analysis([frame], source_names=["orders"])
    assert next(f for f in plan.fact_metrics if f.field == "销售额").unit == unit
    result = build_adaptive_analysis_report([frame], source_names=["orders"])
    assert result.outputs["管理层通用总览"].set_index("指标").loc["核心指标：销售额", "单位"] == unit
    assert result.outputs["语义契约"].iloc[0]["单位"] == unit
    assert result.outputs["分类排名"]["单位"].eq(unit).all()
    assert result.outputs["自适应图表看板"]["排名单位"].eq(unit).all()
    assert result.outputs["自适应图表看板"]["趋势单位"].eq(unit).all()
    pd.testing.assert_frame_equal(original, frame)


@pytest.mark.parametrize("currencies", [
    ["USD", "CNY", "USD", "USD"], ["USD", None, "USD", "USD"],
    ["USD", "", "USD", "USD"], ["UNKNOWN"] * 4,
])
def test_mixed_missing_or_unknown_currency_blocks_money(currencies):
    frame = sales().assign(币种=currencies)
    assert metric_contract(0, "orders", frame, "销售额").status == "unavailable"


def test_currency_aliases_are_not_treated_as_different_currencies():
    frame = sales().assign(币种=["CNY", "RMB", "人民币", " cny "])
    contract = metric_contract(0, "orders", frame, "销售额")
    assert contract.status == "inferred" and contract.unit == "元"


@pytest.mark.parametrize("field,value", [
    ("销售额(元)", 100), ("销售额(万元)", 100), ("销售额", "￥100"), ("销售额", "¥100"),
])
def test_currency_conflict_with_header_or_cell_symbol_is_blocked(field, value):
    frame = sales().rename(columns={"销售额": field})
    frame[field] = [value, 200, 300, 400]
    frame["币种"] = "USD"
    result = build_adaptive_analysis_report([frame], source_names=["orders"])
    assert result.outputs["管理层通用总览"].set_index("指标").loc[f"核心指标：{field}", "结果"] == "不可计算"
    assert result.outputs["主数据分析"].iloc[0][field] == value


def test_ratio_components_require_compatible_currency_and_nonpercent_amounts():
    frame = sales().rename(columns={"销售额": "销售额(USD)"})
    frame["利润(EUR)"] = [10, 20, 30, 40]
    frame["利润率"] = ["10%"] * 4
    assert metric_contract(0, "orders", frame, "利润率").status == "unavailable"
    frame = sales().assign(币种="USD", 利润=[10, 20, 30, 40], 利润率=["10%"] * 4)
    assert evaluate_metric(frame, metric_contract(0, "orders", frame, "利润率"))[0] == pytest.approx(.1)
    frame["利润"] = ["10%", "20", "30", "40"]
    assert metric_contract(0, "orders", frame, "利润率").status == "unavailable"


def test_conflicting_currency_declarations_and_irrelevant_currency():
    frame = sales().assign(币种="USD", currency="EUR")
    assert metric_contract(0, "orders", frame, "销售额").status == "unavailable"
    # This is a different metric's currency, not a table-wide declaration.
    frame = sales().assign(采购币种="USD")
    assert metric_contract(0, "orders", frame, "销售额").unit != "USD"


def test_currency_and_ratio_survive_verified_excel_export(tmp_path):
    from zipfile import ZipFile
    import xml.etree.ElementTree as ET
    from openpyxl import load_workbook
    from excel_data_toolbox.delivery_qa import export_verified
    from excel_data_toolbox.core import export_tables

    frame = sales().assign(币种="USD", 利润=[10, 40, 90, 160], 利润率=["10%", "20%", "30%", "40%"])
    original = frame.copy(deep=True)
    result = build_adaptive_analysis_report([frame], source_names=["orders"])
    path = tmp_path / "synthetic_currency_report.xlsx"
    acceptance = export_verified(result.outputs, path,
        writer=lambda data, dest: export_tables(data, dest, include_log=False, overwrite=True))
    assert acceptance.status == "passed"
    workbook = load_workbook(path, data_only=True)
    try:
        rows = list(workbook["管理层通用总览"].values)
        sales_row = next(row for row in rows if "核心指标：销售额" in row)
        ratio_row = next(row for row in rows if "核心指标：利润率" in row)
        assert sales_row[:3] == ("核心指标：销售额", 1000, "USD")
        assert ratio_row[1] == pytest.approx(.3)
        assert len(workbook["自适应图表看板"]._charts) == 3
    finally:
        workbook.close()
    with ZipFile(path) as archive:
        charts = [ET.fromstring(archive.read(name)) for name in archive.namelist()
                  if name.startswith("xl/charts/chart") and name.endswith(".xml")]
    namespaces = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
                  "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    unit_labels = [node.text for chart in charts for node in chart.findall(".//c:catAx/c:title//a:t", namespaces)]
    unit_labels += [node.text for chart in charts for node in chart.findall(".//c:valAx/c:title//a:t", namespaces)]
    assert unit_labels.count("指标值（USD）") == 2
    pd.testing.assert_frame_equal(original, frame)


def test_invalid_amount_format_is_preserved_in_verified_excel(tmp_path):
    from openpyxl import load_workbook
    from excel_data_toolbox.delivery_qa import export_verified
    from excel_data_toolbox.core import export_tables

    frame = sales().assign(销售额=["10%", "200.5", "300", "400"])
    result = build_adaptive_analysis_report([frame], source_names=["orders"])
    path = tmp_path / "synthetic_invalid_amount.xlsx"
    assert export_verified(result.outputs, path,
        writer=lambda data, dest: export_tables(data, dest, include_log=False, overwrite=True)).status == "passed"
    workbook = load_workbook(path, data_only=True)
    try:
        assert any("10%" in row for row in workbook["主数据分析"].values)
        row = next(row for row in workbook["管理层通用总览"].values if "核心指标：销售额" in row)
        assert row[1] == "不可计算"
        assert len(workbook["自适应图表看板"]._charts) == 1  # raw-row counts only, not money
    finally:
        workbook.close()
