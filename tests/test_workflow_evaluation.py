from dataclasses import replace

import pandas as pd
import pytest
from openpyxl import load_workbook

from excel_data_toolbox.ai_evaluation import EvaluationScenario, run_evaluation
from excel_data_toolbox.nl_agent import AgentPlan, AgentStep
from excel_data_toolbox.workflow_evaluation import OutputCheck, WorkflowScenario, run_workflow_evaluation


def fixture():
    return {"orders": pd.DataFrame({"订单号": ["A", "B", "C", "D"],
        "日期": ["2026-01-01", "2026-01-02", "2026-02-01", "2026-02-02"],
        "地区": ["东", "西", "东", "西"], "销售额": [100, 200, 300, 400], "利润": [30, 40, 80, 100]})}


def mock_planner(prompt, catalog):
    """A contract fixture, not proof that a remote LLM understands paraphrases."""
    ids = [t["table_id"] for t in catalog["tables"]]
    return AgentPlan(1, "ready", "Synthetic workflow acceptance", "", (), (), (), (
        AgentStep("report", "adaptive_analysis_report", tuple(ids), "report",
            {"source_names": ids, "user_request": prompt}),))


def scenario(name="golden_sales", **kwargs):
    return WorkflowScenario(name, "全面分析地区利润、排名和月度趋势，生成老板报表", ("adaptive_analysis_report",),
        (OutputCheck("管理层通用总览", "结果", 250, "指标", "核心指标：利润"),
         OutputCheck("管理层通用总览", "结果", 1000, "指标", "核心指标：销售额")),
        ("语义契约", "问题与证据", "事实域明细"), **kwargs)


def test_execution_and_workbook_are_checked_not_only_operation_names(tmp_path):
    result = run_workflow_evaluation([scenario(minimum_charts=3)], fixture(),
        planner=mock_planner, planner_kind="mock", destination_dir=tmp_path)
    assert result["planner_kind"] == "mock"
    assert result["passed"] == 1 and result["failed"] == 0, result
    assert result["results"][0]["workbook_verified"]
    workbook = load_workbook(tmp_path / "golden_sales.xlsx")
    try:
        assert "语义契约" in workbook.sheetnames
        assert "问题与证据" in workbook.sheetnames
        assert len(workbook["自适应图表看板"]._charts) >= 3
    finally:
        workbook.close()


def test_wrong_business_values_fail_before_publication(tmp_path):
    wrong = replace(scenario(), checks=(OutputCheck("管理层通用总览", "结果", 999, "指标", "核心指标：销售额"),))
    result = run_workflow_evaluation([wrong], fixture(), planner=mock_planner, planner_kind="mock", destination_dir=tmp_path)
    assert result["failed"] == 1
    assert "output_check:1" in result["results"][0]["failures"]
    assert not list(tmp_path.glob("*.xlsx"))


def test_chart_failure_does_not_publish_a_deliverable(tmp_path):
    result = run_workflow_evaluation([scenario(minimum_charts=999)], fixture(),
        planner=mock_planner, planner_kind="mock", destination_dir=tmp_path)
    assert result["failed"] == 1
    assert not list(tmp_path.glob("*.xlsx"))


def test_status_alias_is_compatible_with_real_planner_contract():
    report = run_evaluation([EvaluationScenario("clarify_case", "删除所有异常记录", expected_status="needs_clarification")],
        lambda _: {"status": "clarification", "steps": []})
    assert report.passed == 1


@pytest.mark.parametrize("prompt", [
    "全面分析这些陌生数据，找出指标、趋势、排名和异常，生成老板看的 Excel 报表。",
    "帮我看看整体表现，按月份看变化、按地区看表现，把口径问题单列，输出经营看板。",
    "请整合当前数据，分析地区表现和月度趋势，最后做一份管理层直接看的 Excel。",
])
def test_real_local_route_paraphrases_reach_verified_download(monkeypatch, tmp_path, prompt):
    import excel_data_toolbox.server as server
    monkeypatch.setattr(server, "TASK_REPOSITORY", server.TaskRepository(tmp_path / "tasks"))
    session = server.AppSession()
    monkeypatch.setattr(server, "SESSION", session)
    monkeypatch.setattr(server, "_project_ai_config", lambda: {"configured": False, "api_key": "", "model": "deepseek-v4-flash"})
    def no_network(*_, **__):
        raise AssertionError("Synthetic local acceptance must not call the API")
    monkeypatch.setattr(server.DeepSeekClient, "classify_unified_request", no_network)
    try:
        raw = fixture()["orders"].rename(columns={"销售额": "合同金额", "利润": "服务利润", "地区": "运营区域"})
        original = raw.copy(deep=True)
        table_id = session.add_table("服务记录", raw, source="synthetic", original=True)
        handler = object.__new__(server.ToolboxHandler)
        planned = handler._ai_unified({"prompt": prompt, "table_ids": [table_id]})
        assert planned["plan"]["steps"][0]["operation"] == "adaptive_analysis_report"
        executed = handler._ai_execute({"plan_token": planned["plan_token"], "confirmed": True})
        assert executed["steps_completed"] == 1
        assert executed["download_url"]
        exported = next(session.output_dir.glob("*.xlsx"))
        book = load_workbook(exported)
        try:
            assert {"语义契约", "问题与证据", "事实域明细", "数据源确认"}.issubset(book.sheetnames)
            assert len(book["自适应图表看板"]._charts) >= 3
        finally:
            book.close()
        overview = next(t.frame for t in session.tables.values() if t.name == "管理层通用总览")
        assert overview.set_index("指标").loc["核心指标：服务利润", "结果"] == 250
        pd.testing.assert_frame_equal(original, session.get(table_id).frame)
    finally:
        session.close()
