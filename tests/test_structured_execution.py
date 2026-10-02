"""Synthetic end-to-end contracts, not another customer-specific template."""
from dataclasses import replace

import pandas as pd
import pytest
from openpyxl import load_workbook

from excel_data_toolbox.core import join_tables
from excel_data_toolbox.delivery_qa import export_verified, verify_delivery
from excel_data_toolbox import export_tables
from excel_data_toolbox.nl_agent import (
    AgentPlan, AgentStep, AgentExecutionError, PlanValidationError,
    build_table_catalog, execute_plan, validate_plan,
)
from excel_data_toolbox.operation_contracts import ALLOWED_OPERATIONS, PARAM_KEYS
from excel_data_toolbox.tool_registry import build_builtin_registry


def plan(*steps):
    return AgentPlan(1, "ready", "Explicit structured processing", "", (), (), (), tuple(steps))


def step(name, operation, inputs, params):
    return AgentStep(name, operation, tuple(inputs), name, params)


def test_pipeline_preserves_exclusions_and_source_and_audits_every_step():
    raw = pd.DataFrame({"id": ["Ａ1 ", "A1", "A2", "A3"], "store": ["East"] * 4,
                        "status": ["paid", "paid", "cancelled", "paid"],
                        "revenue": [100, 100, 500, 90], "cost": [60, 60, 400, 20]})
    original = raw.copy(deep=True)
    workflow = plan(
        step("normalize", "normalize_fields", ["raw"], {"fields": {"id": {"case": "upper"}}}),
        step("dedup", "deduplicate", ["$normalize"], {"subset": ["id"]}),
        step("valid", "partition", ["$dedup"], {"rules": [{"conditions": [
            {"column": "status", "operator": "ne", "value": "paid"}], "reason": "Not a paid order"}]}),
        step("compute", "derive_columns", ["$valid"], {"formulas": [
            {"output": "profit", "operator": "subtract", "left": {"column": "revenue"}, "right": {"column": "cost"}}]}),
        step("totals", "summary", ["$compute"], {"by": "store", "aggregations": {"revenue": "sum", "profit": "sum"}}),
    )
    result = execute_plan(workflow, {"raw": raw}, dry_run=False)
    pd.testing.assert_frame_equal(original, raw)
    assert result.tables["totals"].iloc[0]["revenue"] == 190
    assert result.tables["totals"].iloc[0]["profit"] == 110
    assert len(result.tables["dedup_rejected"]) == 1
    assert len(result.tables["valid_rejected"]) == 1
    assert result.tables["valid_rejected"]["排除原因"].iloc[0] == "Not a paid order"
    for report in result.reports.values():
        audit = report["execution"]
        assert audit["status"] == "completed"
        assert len(audit["parameter_sha256"]) == 64
        assert len(audit["inputs"][0]["fingerprint"]) == 64
        assert audit["duration_ms"] >= 0
        assert "paid" not in str(audit)


def test_registry_and_execution_share_exact_contracts():
    registry = build_builtin_registry(sorted(ALLOWED_OPERATIONS))
    assert registry.names() == ALLOWED_OPERATIONS
    for name in registry.names():
        tool = registry.get(name)
        assert tool.executor is not None and tool.validator is not None
        required, allowed = PARAM_KEYS[name]
        assert set(tool.definition.parameter_schema["properties"]) == allowed
        assert set(tool.definition.parameter_schema["required"]) == required
    with pytest.raises(ValueError, match="缺少参数"):
        registry.execute("summary", {}, [pd.DataFrame({"x": [1]})])
    with pytest.raises(ValueError, match="未知参数"):
        registry.execute("clean", {"arbitrary": True}, [pd.DataFrame()])


@pytest.mark.parametrize("bad", [[], ["raw", "extra"]])
def test_direct_dataclass_cannot_bypass_input_arity(bad):
    with pytest.raises(PlanValidationError):
        execute_plan(plan(step("s", "clean", bad, {})), {"raw": pd.DataFrame()})


def test_direct_plan_cannot_bypass_required_parameters():
    with pytest.raises(PlanValidationError):
        execute_plan(plan(step("s", "summary", ["raw"], {})), {"raw": pd.DataFrame({"x": [1]})})


def test_source_ids_are_not_step_artifact_references():
    with pytest.raises(PlanValidationError, match="上游"):
        execute_plan(plan(step("s", "clean", ["$raw"], {})), {"raw": pd.DataFrame({"x": [1]})})


def test_conflicting_and_missing_duplicate_keys_never_silently_disappear():
    for raw in (pd.DataFrame({"id": ["A", "A"], "amount": [10, 99]}),
                pd.DataFrame({"id": [None, None], "amount": [10, 10]})):
        with pytest.raises(AgentExecutionError):
            execute_plan(plan(step("s", "deduplicate", ["raw"], {"subset": ["id"]})), {"raw": raw})


def test_many_to_many_null_keys_and_row_explosion_are_blocked():
    left = pd.DataFrame({"key": ["A", "A"]})
    right = pd.DataFrame({"key": ["A", "A", "A"]})
    with pytest.raises(ValueError, match="多对多"):
        join_tables(left, right, on="key")
    with pytest.raises(ValueError, match="上限"):
        join_tables(left, right, on="key", allow_many_to_many=True, max_output_rows=5)
    assert len(join_tables(left, right, on="key", allow_many_to_many=True, max_output_rows=6)) == 6
    with pytest.raises(ValueError, match="缺失"):
        join_tables(pd.DataFrame({"key": [None]}), pd.DataFrame({"key": [None]}), on="key")


def test_missing_arithmetic_and_zero_denominator_are_not_zero_profit():
    raw = pd.DataFrame({"sales": [10.0, 0.0, None], "profit": [2.0, 3.0, 1.0]})
    params = {"formulas": [{"output": "margin", "operator": "divide", "left": {"column": "profit"},
                            "right": {"column": "sales"}}]}
    with pytest.raises(AgentExecutionError):
        execute_plan(plan(step("s", "derive_columns", ["raw"], params)), {"raw": raw})
    result = execute_plan(plan(step("s", "derive_columns", ["raw"], {**params, "errors": "coerce"})), {"raw": raw})
    assert result.tables["s"]["margin"].iloc[0] == 0.2
    assert result.tables["s"]["margin"].iloc[1:].isna().all()
    assert len(result.tables["s_review"]) == 2


def test_quality_gate_stops_chain_and_success_passes_through_input():
    rules = [{"rule_id": "required", "rule_type": "not_null", "column": "amount"}]
    workflow = plan(step("gate", "validate", ["raw"], {"rules": rules, "fail_on_error": True}),
                    step("after", "clean", ["$gate"], {}))
    with pytest.raises(AgentExecutionError, match="门禁"):
        execute_plan(workflow, {"raw": pd.DataFrame({"amount": [None]})})
    result = execute_plan(workflow, {"raw": pd.DataFrame({"amount": [2.0]})})
    assert result.tables["after"].iloc[0]["amount"] == 2


def test_large_fact_catalog_is_not_limited_to_twenty_inputs():
    tables = {f"table_{i}": pd.DataFrame({"x": [i]}) for i in range(25)}
    workflow = plan(step("all", "adaptive_analysis_report", list(tables), {"source_names": list(tables)}))
    assert validate_plan(workflow.to_dict(), build_table_catalog(tables)).executable


def writer(frames, path):
    export_tables(frames, path, include_log=False, overwrite=True)


def test_delivery_preserves_literal_na_and_rejects_undeclared_sheets(tmp_path):
    tables = {"raw": pd.DataFrame({"id": ["NA", "N/A", "001"], "amount": [2, 4, 6]})}
    path = tmp_path / "result.xlsx"
    report = export_verified(tables, path, writer=writer)
    assert report.status == "passed" and report.artifact == path.name
    workbook = load_workbook(path)
    workbook.create_sheet("surprise")
    workbook.save(path)
    workbook.close()
    assert verify_delivery(path, tables, allow_extra_tables=False).status == "failed"
    assert verify_delivery(path, tables, allow_extra_tables=True).status == "passed"


@pytest.mark.parametrize("failure", ["writer", "acceptance", "domain"])
def test_failed_delivery_preserves_previous_file_and_removes_staging(tmp_path, failure):
    path = tmp_path / "result.xlsx"
    tables = {"raw": pd.DataFrame({"amount": [10]})}
    export_verified(tables, path, writer=writer)
    previous = path.read_bytes()
    def failing_writer(frames, target):
        if failure == "writer":
            raise OSError("simulated writer failure")
        writer({"wrong": pd.DataFrame({"amount": [99]})} if failure == "acceptance" else frames, target)
    def failing_domain(_):
        raise ValueError("domain check failure")
    with pytest.raises((ValueError, OSError)):
        export_verified(tables, path, writer=failing_writer,
                        extra_validator=failing_domain if failure == "domain" else None)
    assert path.read_bytes() == previous
    assert not list(tmp_path.glob(".delivery-*"))


def test_empty_acceptance_and_appended_records_cannot_pass(tmp_path):
    path = tmp_path / "result.xlsx"
    tables = {"raw": pd.DataFrame({"amount": [10]})}
    writer(tables, path)
    with pytest.raises(ValueError):
        verify_delivery(path, {})
    workbook = load_workbook(path)
    workbook["raw"].append([99])
    workbook.save(path)
    workbook.close()
    assert verify_delivery(path, tables).status == "failed"


def test_failed_server_delivery_never_commits_success(monkeypatch, tmp_path):
    import excel_data_toolbox.server as server
    monkeypatch.setattr(server, "TASK_REPOSITORY", server.TaskRepository(tmp_path / "tasks"))
    session = server.AppSession()
    monkeypatch.setattr(server, "SESSION", session)
    table_id = session.add_table("sales", pd.DataFrame({"amount": [10]}), source="synthetic", original=True)
    workflow = plan(step("report", "adaptive_analysis_report", [table_id], {"source_names": ["sales"]}))
    ticket = session.issue_ai_plan(table_ids=[table_id],
                                  table_signatures=(server._ai_table_signature(session.get(table_id)),),
                                  plan=workflow, model="no-network-test")
    generated = {name: pd.DataFrame({"amount": [10]}) for name in (
        "管理层通用总览", "主数据分析", "数据字典", "数据质量", "表关系建议", "分类排名", "时间趋势", "异常数据", "自适应图表看板")}
    real_result = execute_plan(plan(step("clean", "clean", [table_id], {})), {table_id: session.get(table_id).frame})
    monkeypatch.setattr(server, "execute_plan", lambda *_args, **_kwargs: replace(real_result, tables=generated))
    def fail(*_, **__):
        raise ValueError("forced delivery rejection")
    monkeypatch.setattr(server, "export_verified", fail)
    handler = object.__new__(server.ToolboxHandler)
    with pytest.raises(server.ApiError, match="forced delivery rejection"):
        handler._ai_execute({"plan_token": ticket, "confirmed": True})
    assert set(session.tables) == {table_id}
    assert session.operations == [] and session.review_items == {} and session.downloads == {}
    assert session.active_table == table_id


def test_failed_session_persistence_rolls_back_result_and_success_record(monkeypatch, tmp_path):
    import excel_data_toolbox.server as server
    monkeypatch.setattr(server, "TASK_REPOSITORY", server.TaskRepository(tmp_path / "tasks"))
    session = server.AppSession()
    monkeypatch.setattr(server, "SESSION", session)
    table_id = session.add_table("raw", pd.DataFrame({"value": [1]}), source="synthetic", original=True)
    workflow = plan(step("clean", "clean", [table_id], {}))
    ticket = session.issue_ai_plan(table_ids=[table_id],
                                  table_signatures=(server._ai_table_signature(session.get(table_id)),),
                                  plan=workflow, model="no-network-test")
    prior_redo = {"marker": "preserve redo history"}
    session.redo_stack.append(prior_redo)
    def fail():
        raise OSError("simulated disk failure")
    monkeypatch.setattr(session, "persist", fail)
    with pytest.raises(OSError, match="disk failure"):
        object.__new__(server.ToolboxHandler)._ai_execute({"plan_token": ticket, "confirmed": True})
    assert set(session.tables) == {table_id}
    assert session.operations == [] and session.history == [] and session.downloads == {}
    assert session.redo_stack == [prior_redo] and session.active_table == table_id
