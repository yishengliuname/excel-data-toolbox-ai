"""Rule reuse must reject drift and missing authority, not memorize answers."""
import copy
import json

import pytest

from excel_data_toolbox.autopilot import run_delivery
from excel_data_toolbox.business_rules import BusinessRuleError
from excel_data_toolbox.business_workflow import BusinessProfileStore
from excel_data_toolbox.task_acceptance import assess_task
from excel_data_toolbox.trusted_automation import TrustedAutomationStore, validate_limits
from tests.test_business_rules_workflow import data, rules, setup_server


REQUEST = '分析月度毛利与门店排名'
LIMITS = {'valid_days': 30, 'max_rows': 100, 'max_excluded_ratio': .8}


def grant(tmp_path):
    store = TrustedAutomationStore(tmp_path)
    profile = store.save_authorized('A', 'monthly', rules(), [data()], ['销售明细'], REQUEST, LIMITS)
    return store, profile['profile_id']


def test_old_profile_is_review_only_and_reuse_does_not_read_customer_rows(tmp_path):
    saved = BusinessProfileStore(tmp_path).save('A', 'old', rules(), [data()], ['销售明细'])
    store = TrustedAutomationStore(tmp_path)
    with pytest.raises(BusinessRuleError, match='授权'):
        store.authorized_rules('A', saved['profile_id'], [data()], ['销售明细'], REQUEST)
    store, identifier = grant(tmp_path)
    fresh = data()
    fresh.loc[[0, 1], 'Q'] = '150'  # fresh source values, not cached totals
    policy, proof = store.authorized_rules('A', identifier, [fresh], ['新文件__销售明细'], REQUEST)
    assert policy['tables'][0]['source_name'] == '新文件__销售明细'
    assert proof['source_checks'][0]['included_rows'] == 2
    assert '2026-08-01' not in next(p for p in tmp_path.rglob('*.json') if identifier in p.name).read_text(encoding='utf-8')
    assert store.list('A')[-1]['automation_enabled'] is True or any(p['automation_enabled'] for p in store.list('A'))


@pytest.mark.parametrize('change,match', [
    ('customer', '客户'), ('status', '状态'), ('conflict', '冲突'), ('schema', '漂移'),
    ('date', '日期'), ('unit', '币种|单位'), ('rows', '记录数'), ('excluded', '排除比例'),
    ('topic', '主题'), ('revoked', '撤销'), ('expired', '过期'), ('engine', '引擎'), ('rules', '结构'),
])
def test_authorization_drift_blocks_only_with_specific_reason(tmp_path, change, match):
    store, identifier = grant(tmp_path)
    frame, customer, request = data(), 'A', REQUEST
    path = next(tmp_path.rglob(f'{identifier}.json'))
    record = json.loads(path.read_text(encoding='utf-8'))
    if change == 'customer':
        customer = 'B'
    elif change == 'status':
        frame.loc[2, '原始状态'] = '未知回调状态'
    elif change == 'conflict':
        frame.loc[1, 'Z'] = 999
    elif change == 'schema':
        frame = frame.rename(columns={'Q': 'Q2'})
    elif change == 'date':
        frame.loc[2, '业务日期'] = 'not-a-date'
    elif change == 'unit':
        # Drift within a pre-existing currency declaration schema is also
        # tested by business rules; new currency fields block schema first.
        frame = frame.assign(Q币种='USD')
        record['schemas']['销售明细'] = list(frame.columns)
        from excel_data_toolbox.trusted_automation import _digest
        record['automation']['contract_digest'] = _digest({'rules': record['rules'], 'schemas': record['schemas']})
    elif change == 'rows':
        record['automation']['limits']['max_rows'] = 5
    elif change == 'excluded':
        record['automation']['limits']['max_excluded_ratio'] = .1
    elif change == 'topic':
        request = '分析库存是否需要补货以及积压'
    elif change == 'revoked':
        store.revoke('A', identifier)
        record = json.loads(path.read_text(encoding='utf-8'))
    elif change == 'expired':
        record['automation']['expires_at'] = '2020-01-01T00:00:00+00:00'
    elif change == 'engine':
        record['automation']['engine_revision'] = 'old'
    elif change == 'rules':
        record['rules']['tables'][0]['date_field'] = '对账日期'
    path.write_text(json.dumps(record, ensure_ascii=False), encoding='utf-8')
    with pytest.raises(BusinessRuleError, match=match):
        store.authorized_rules(customer, identifier, [frame], ['销售明细'], request)


@pytest.mark.parametrize('limits', [True, {}, {**LIMITS, 'valid_days': True}, {**LIMITS, 'max_rows': 0},
                                    {**LIMITS, 'max_excluded_ratio': float('nan')}, {**LIMITS, 'exec': True}])
def test_authorization_limits_are_explicit_and_bounded(limits):
    with pytest.raises(BusinessRuleError):
        validate_limits(limits)


def test_revocation_between_planning_and_execution_cannot_reuse_revision(tmp_path):
    store, identifier = grant(tmp_path)
    _, proof = store.authorized_rules('A', identifier, [data()], ['销售明细'], REQUEST)
    store.revoke('A', identifier)
    with pytest.raises(BusinessRuleError):
        with store.execution_guard('A', identifier, [data()], ['销售明细'], REQUEST, revision=proof['revision']):
            pytest.fail('must not execute')


@pytest.mark.parametrize('field,value', [('expires_at', None), ('expires_at', '2027-01-01'), ('topics', None)])
def test_malformed_grant_is_a_specific_review_error_not_unhandled_exception(tmp_path, field, value):
    store, identifier = grant(tmp_path)
    path = next(tmp_path.rglob(f'{identifier}.json'))
    record = json.loads(path.read_text(encoding='utf-8'))
    record['automation'][field] = value
    path.write_text(json.dumps(record, ensure_ascii=False), encoding='utf-8')
    with pytest.raises(BusinessRuleError, match='结构无效'):
        store.authorized_rules('A', identifier, [data()], ['销售明细'], REQUEST)


def test_task_acceptance_does_not_confuse_file_success_with_goal_success():
    result = {'status': 'completed', 'steps_completed': 1, 'inputs_unchanged': True, 'pending_reviews': 0,
              'download_url': '/download/' + 'a' * 32, 'acceptance': {'status': 'passed'},
              'request_coverage': [{'topic': 'trend', 'status': 'available'}]}
    assert assess_task(result, expected_steps=1, required_topics=['trend'])['status'] == 'passed'
    assert assess_task(result, expected_steps=1, required_topics=['trend', 'cash'])['status'] == 'needs_review'
    for field in ('inputs_unchanged', 'acceptance', 'download_url', 'steps_completed'):
        missing = copy.deepcopy(result)
        missing.pop(field)
        assert assess_task(missing, expected_steps=1)['status'] != 'passed'


def test_transient_planning_retry_is_bounded_and_execution_never_retries():
    calls = []
    def plan():
        calls.append('plan')
        if len(calls) == 1:
            raise TimeoutError('private URL deliberately not stored')
        return {'mode': 'data', 'status': 'ready', 'auto_execute': True, 'plan_token': 't'}
    result = run_delivery(plan=plan, execute=lambda _: {'status': 'completed', 'download_url': '/download/' + 'a'*32},
                          export=lambda _: {}, progress=lambda *_: None)
    assert calls == ['plan', 'plan'] and result['planning_recovery']['attempts'] == 2
    assert result['planning_recovery']['events'] == [{'attempt': 1, 'reason': 'TimeoutError'}]
    calls.clear()
    with pytest.raises(TimeoutError):
        run_delivery(plan=lambda: {'mode': 'data', 'status': 'ready', 'auto_execute': True, 'plan_token': 't'},
                     execute=lambda _: (calls.append('execution'), (_ for _ in ()).throw(TimeoutError()))[1],
                     export=lambda _: {}, progress=lambda *_: None)
    assert calls == ['execution']
    attempts = []
    def always_fail():
        attempts.append(1)
        raise ConnectionError()
    with pytest.raises(ConnectionError):
        run_delivery(plan=always_fail, execute=lambda _: pytest.fail(), export=lambda _: {}, progress=lambda *_: None)
    assert len(attempts) == 2


def test_human_confirmation_can_opt_in_but_unconfirmed_payload_cannot(monkeypatch, tmp_path):
    server, session, handler, table = setup_server(monkeypatch, tmp_path)
    draft = handler._business_review({'prompt': REQUEST, 'table_ids': [table]})
    payload = {'review_token': draft['review_token'], 'rules': rules(), 'customer_scope': 'A',
               'profile_title': 'monthly', 'automation_limits': LIMITS}
    with pytest.raises(server.ApiError, match='人工'):
        handler._business_confirm({**payload, 'confirmed': False})
    plan = handler._business_confirm({**payload, 'confirmed': True})
    assert plan['saved_profile']['automation']['enabled']
    identifier = plan['saved_profile']['profile_id']
    assert handler._business_profiles({'customer_scope': 'A'})['profiles'][0]['automation_enabled']
    handler._business_profile_revoke({'customer_scope': 'A', 'profile_id': identifier, 'confirmed': True})
    assert not handler._business_profiles({'customer_scope': 'A'})['profiles'][0]['automation_enabled']


def test_http_authorized_new_data_exports_new_numbers_and_revocation_blocks_cache(monkeypatch, tmp_path):
    import hashlib
    import http.client
    import io
    import threading
    import time
    from openpyxl import load_workbook
    from excel_data_toolbox.session_registry import SessionProxy, SessionRegistry
    from excel_data_toolbox.task_engine import PersistentJobEngine

    server, session, _, table = setup_server(monkeypatch, tmp_path)
    registry = SessionRegistry(lambda i: server.AppSession(restore_task_id=i), session)
    monkeypatch.setattr(server, 'SESSION_REGISTRY', registry)
    monkeypatch.setattr(server, 'SESSION', SessionProxy(registry))
    engine = PersistentJobEngine(tmp_path / 'jobs.sqlite', workers=1)
    engine.register('business_autopilot', server._job_business_autopilot)
    monkeypatch.setattr(server, 'JOB_ENGINE', engine)
    # The authorized path must not need or silently call an AI provider.
    monkeypatch.setattr(server.ToolboxHandler, '_ai_unified', lambda *a, **k: pytest.fail('authorized rules are local'))
    httpd = server.ThreadingHTTPServer(('127.0.0.1', 0), server.ToolboxHandler)
    thread = threading.Thread(target=httpd.serve_forever, daemon=True)
    thread.start()
    connection = http.client.HTTPConnection('127.0.0.1', httpd.server_address[1], timeout=30)
    def request(path, payload=None, expected=200):
        connection.request('GET' if payload is None else 'POST', path,
            None if payload is None else json.dumps(payload, ensure_ascii=False).encode(),
            {'Content-Type': 'application/json', 'X-Task-ID': session.task_id})
        response = connection.getresponse()
        contents = response.read()
        assert response.status == expected, contents.decode(errors='replace')
        return json.loads(contents) if 'application/json' in response.headers.get('Content-Type', '') else contents
    engine.start()
    try:
        draft = request('/api/business-review', {'prompt': REQUEST, 'table_ids': [table]})
        plan = request('/api/business-confirm', {'review_token': draft['review_token'], 'rules': rules(),
            'confirmed': True, 'customer_scope': 'A', 'profile_title': 'monthly', 'automation_limits': LIMITS})
        identifier = plan['saved_profile']['profile_id']
        with session.lock:
            session.get(table).frame.loc[[0, 1], 'Q'] = '150'
            session.get(table).frame.loc[2, 'Q'] = '25'
            session.persist()
        before = hashlib.sha256(session.get(table).frame.to_json().encode()).hexdigest()
        goal = {'prompt': REQUEST, 'table_ids': [table], 'customer_scope': 'A', 'profile_id': identifier}
        job = request('/api/ai-run', goal)['job']
        deadline = time.monotonic() + 25
        while job['status'] in {'queued', 'running'} and time.monotonic() < deadline:
            time.sleep(.05)
            job = request('/api/jobs')['jobs'][0]
        assert job['status'] == 'completed', job
        result = job['result']
        assert result['authorization']['profile_id'] == identifier
        assert result['task_acceptance']['status'] == 'needs_review'
        assert result['delivery_status'] == 'partial'  # gross contribution != net profit
        assert result['execution']['acceptance']['status'] == 'passed'
        assert result['execution']['inputs_unchanged']
        assert all(c['status'] == 'passed' for c in result['task_acceptance']['checks'][:4])
        workbook = load_workbook(io.BytesIO(request(result['download_url'])))
        try:
            metrics = {r[0]: r[1] for r in workbook['管理层通用总览'].values if len(r) > 1}
            assert metrics['核心指标：Q'] == 175  # changed data, not original 120
            assert metrics['核心指标：毛利'] == 130
            excluded_rows = list(workbook['规则排除明细'].values)
            header_index = next(i for i, row in enumerate(excluded_rows) if '排除源行号' in row)
            assert len([row for row in excluded_rows[header_index + 1:] if any(v is not None for v in row)]) == 4
            assert len(workbook['自适应图表看板']._charts) >= 2
            categories = [row[0] for row in list(workbook['分类排名'].values)[4:] if row[0] is not None]
            assert '2026-09-01' not in categories
        finally:
            workbook.close()
        assert before == hashlib.sha256(session.get(table).frame.to_json().encode()).hexdigest()
        same = request('/api/ai-run', goal)['job']
        assert same['job_id'] == job['job_id']
        request('/api/business-profile-revoke', {'customer_scope': 'A', 'profile_id': identifier, 'confirmed': True})
        assert '撤销' in request('/api/ai-run', goal, expected=422)['error']
    finally:
        connection.close()
        httpd.shutdown()
        httpd.server_close()
        thread.join(timeout=5)
        engine.stop(timeout=5)


def test_trusted_selected_dimension_can_supply_evidence_without_amount_inflation(tmp_path):
    import pandas as pd
    facts = pd.DataFrame({'商品编码': ['A', 'A', 'B'], '销售额': [10, 20, 30]})
    master = pd.DataFrame({'SKU': ['A', 'B'], '类别': ['food', 'daily']})
    policy = {'schema_version': 1, 'tables': [
        {'source_name': '订单', 'role': 'fact', 'metrics': {'销售额': {'aggregation': 'sum', 'unit': '元'}}},
        {'source_name': '商品', 'role': 'dimension'}],
        'joins': [{'left': '订单', 'right': '商品', 'left_key': '商品编码', 'right_key': 'SKU', 'fields': ['类别']}]}
    store = TrustedAutomationStore(tmp_path)
    identifier = store.save_authorized('B', '分类经营', policy, [facts, master], ['订单', '商品'], '分析商品销售额排名', LIMITS)['profile_id']
    _, proof = store.authorized_rules('B', identifier, [facts, master], ['订单', '商品'], '分析商品销售额排名')
    assert proof['source_checks'][0]['confirmed_joins'][0]['right'] == '商品'
    assert proof['source_checks'][0]['input_rows'] == proof['source_checks'][0]['included_rows'] == 3
    with pytest.raises(BusinessRuleError, match='覆盖'):
        store.authorized_rules('B', identifier, [facts, master.iloc[:1]], ['订单', '商品'], '分析商品销售额排名')


def test_same_schema_currency_values_cannot_override_confirmed_unit(tmp_path):
    frame = data().assign(Q币种='CNY')
    store = TrustedAutomationStore(tmp_path)
    identifier = store.save_authorized('A', 'currency', rules(), [frame], ['销售明细'], REQUEST, LIMITS)['profile_id']
    frame.loc[2, 'Q币种'] = 'USD'
    with pytest.raises(BusinessRuleError, match='币种|单位'):
        store.authorized_rules('A', identifier, [frame], ['销售明细'], REQUEST)


def test_secondary_dates_are_context_not_operating_ranking_dimensions():
    from excel_data_toolbox.analysis_compiler import compile_analysis
    from excel_data_toolbox.business_rules import prepare_business_data
    prepared = prepare_business_data([data()], ['销售明细'], rules())
    plan = compile_analysis(prepared.frames, source_names=['销售明细'], user_request=REQUEST,
                            business_contexts=prepared.contexts)
    assert plan.dates == ('业务日期',)
    assert '对账日期' not in plan.dimensions
    assert [f.role for f in plan.fields if f.field == '对账日期'] == ['context_date']
    assert all(s.dimension != '对账日期' for s in plan.analyses)
    assert any(s.dimension == '门店' and s.kind == 'ranking' for s in plan.analyses)


def test_queued_job_rechecks_revocation_before_any_business_execution(monkeypatch, tmp_path):
    from types import SimpleNamespace
    from excel_data_toolbox.session_registry import SessionProxy, SessionRegistry
    from excel_data_toolbox.task_engine import PersistentJobEngine
    server, session, handler, table = setup_server(monkeypatch, tmp_path)
    registry = SessionRegistry(lambda i: server.AppSession(restore_task_id=i), session)
    monkeypatch.setattr(server, 'SESSION_REGISTRY', registry)
    monkeypatch.setattr(server, 'SESSION', SessionProxy(registry))
    engine = PersistentJobEngine(tmp_path / 'jobs.sqlite')
    engine.register('business_autopilot', server._job_business_autopilot)
    monkeypatch.setattr(server, 'JOB_ENGINE', engine)
    store, identifier = grant(server.TASK_REPOSITORY.root / '.business_profiles')
    with registry.scoped(session.task_id):
        job = handler._autonomous_submit({'prompt': REQUEST, 'table_ids': [table],
                                         'customer_scope': 'A', 'profile_id': identifier})['job']
    store.revoke('A', identifier)
    monkeypatch.setattr(server.ToolboxHandler, '_ai_execute', lambda *a, **k: pytest.fail('revoked queued grant'))
    before = len(session.tables)
    context = SimpleNamespace(task_id=session.task_id, job_id=job['job_id'], update=lambda *_: None, raise_if_cancelled=lambda: None)
    with pytest.raises(BusinessRuleError, match='撤销'):
        server._job_business_autopilot(context, engine._claim(job['job_id'])[1])
    assert not session.ai_plans and len(session.tables) == before and not session.downloads
