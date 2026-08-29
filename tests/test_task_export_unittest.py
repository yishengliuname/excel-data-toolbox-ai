from __future__ import annotations

import tempfile
from pathlib import Path
import unittest
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

from excel_data_toolbox.core import export_tables, profile_dataframe
from excel_data_toolbox.nl_agent import (
    _normalise_engineering_scope,
    _validate_engineering_brief,
)
from excel_data_toolbox.server import (
    AppSession,
    TableEntry,
    ToolboxHandler,
    _blank_cell_detail_frame,
    _long_text_detail_frame,
)
import excel_data_toolbox.server as server_module


class TaskIsolationAndExportTests(unittest.TestCase):
    def test_reset_handler_returns_new_empty_task_id_for_browser_handoff(self) -> None:
        session = AppSession()
        handler = object.__new__(ToolboxHandler)
        try:
            old_task_id = session.task_id
            session.add_table("旧任务表", pd.DataFrame({"值": [1]}), source="测试")
            with patch.object(server_module, "SESSION", session):
                response = handler._reset({})

            self.assertNotEqual(response["task_id"], old_task_id)
            self.assertEqual(response["task_id"], session.task_id)
            self.assertEqual(response["task_name"], "新建数据处理任务")
            self.assertTrue(response["new_task"])
            self.assertEqual(session.tables, {})
            self.assertEqual(session.state_payload()["tables"], [])
        finally:
            session.close()

    def test_reset_creates_isolated_named_task_folder(self) -> None:
        session = AppSession()
        try:
            first_id = session.task_id
            first_dir = session.task_dir
            session.add_table("旧任务表", pd.DataFrame({"值": [1]}), source="测试")

            session.reset()

            self.assertNotEqual(session.task_id, first_id)
            self.assertEqual(session.task_dir.name, session.task_id)
            self.assertEqual(session.upload_dir.parent, session.task_dir)
            self.assertEqual(session.output_dir.parent, session.task_dir)
            # V9 keeps one durable folder per task until the retention policy
            # or an explicit delete removes it.
            self.assertTrue(first_dir.exists())
            self.assertTrue((first_dir / "manifest.json").exists())
            self.assertEqual(session.tables, {})
        finally:
            session.close()

    def test_export_wraps_long_text_and_preserves_values(self) -> None:
        long_text = "摘要偏少；目录格式错误；图中文字太小；结果分析需要进一步展开。" * 5
        frame = pd.DataFrame({"序号": [1, 2], "问题说明": [long_text, "正常"]})
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "readable.xlsx"
            export_tables({"成绩问题": frame}, path, include_log=False)

            exported = pd.read_excel(path, sheet_name="成绩问题")
            self.assertEqual(exported.shape, frame.shape)
            self.assertEqual(exported.loc[0, "问题说明"], long_text)
            workbook = load_workbook(path)
            worksheet = workbook["成绩问题"]
            self.assertTrue(worksheet["B2"].alignment.wrap_text)
            self.assertGreaterEqual(worksheet.column_dimensions["B"].width, 36)
            self.assertLessEqual(worksheet.column_dimensions["B"].width, 58)
            self.assertGreater(worksheet.row_dimensions[2].height, 60)

    def test_long_text_detail_expands_items_without_mutating_source(self) -> None:
        frame = pd.DataFrame(
            {"序号": [1], "问题说明": ["摘要偏少；目录错误；图中文字太小。"], "得分": [70]}
        )
        original = frame.copy(deep=True)

        detail = _long_text_detail_frame(
            [TableEntry("table-1", "成绩表", frame, "测试", True)]
        )

        self.assertEqual(detail["内容"].tolist(), ["摘要偏少；", "目录错误；", "图中文字太小。"])
        self.assertEqual(detail["记录标识"].tolist(), [1, 1, 1])
        pd.testing.assert_frame_equal(frame, original)

    def test_blank_cell_detail_counts_null_and_whitespace_without_mutation(self) -> None:
        frame = pd.DataFrame(
            {"序号": [1, 2], "得分": [None, 90], "备注": ["  ", "完整"]}
        )
        original = frame.copy(deep=True)

        detail = _blank_cell_detail_frame(
            [TableEntry("table-1", "成绩表", frame, "测试", True)]
        )

        self.assertEqual(len(detail), 2)
        self.assertEqual(set(detail["空值字段"]), {"得分", "备注"})
        self.assertTrue((detail["说明"] == "源数据为空，导出时保持为空").all())
        pd.testing.assert_frame_equal(frame, original)

    def test_profile_counts_empty_strings_as_missing(self) -> None:
        frame = pd.DataFrame({"得分": ["", "  ", 90], "备注": [None, "完整", "完整"]})

        profile = profile_dataframe(frame)

        self.assertEqual(profile.missing_cell_count, 3)
        self.assertEqual(profile.columns[0].missing_count, 2)
        self.assertEqual(profile.columns[0].non_null_count, 1)

    def test_engineering_scope_normalises_safe_model_variants(self) -> None:
        self.assertEqual(
            _normalise_engineering_scope(["只读分析", "不执行外部操作"]),
            "只读分析；不执行外部操作",
        )
        self.assertEqual(
            _normalise_engineering_scope(
                {"包含": ["数据模型", "DAX 指标"], "不包含": "自动发布"}
            ),
            "包含：数据模型、DAX 指标；不包含：自动发布",
        )

    def test_engineering_brief_normalises_object_items_across_all_text_lists(self) -> None:
        payload = {
            "schema_version": 1,
            "status": "ready",
            "category": "power_bi",
            "normalized_request": "生成经营分析工程方案",
            "scope": {"包含": ["星型模型", "DAX"], "不包含": "自动发布"},
            "clarification_questions": [],
            "deliverables": [
                {"name": "经营看板", "description": "包含销售和利润指标"}
            ],
            "implementation_steps": [
                {"step": 1, "action": "建立日期维表", "owner": "人工确认"}
            ],
            "artifacts": [],
            "test_checklist": [{"check": "汇总值一致", "expected": "通过"}],
            "risks": {"risk": "业务口径未确认", "level": "中"},
            "human_approval_points": [{"stage": "发布前", "action": "人工审批"}],
        }

        result = _validate_engineering_brief(payload, "power_bi")

        self.assertEqual(result["status"], "ready")
        self.assertIn("name：经营看板", result["deliverables"][0])
        self.assertIn("action：建立日期维表", result["implementation_steps"][0])
        self.assertIn("check：汇总值一致", result["test_checklist"][0])
        self.assertIn("risk：业务口径未确认", result["risks"][0])

    def test_engineering_deliverable_accepts_detailed_non_executable_content(self) -> None:
        detailed_content = "详细交付说明：" + "经营指标口径、页面结构和验收规则；" * 80
        payload = {
            "schema_version": 1,
            "status": "ready",
            "category": "power_bi",
            "normalized_request": "生成详细经营分析工程方案",
            "scope": "仅生成可人工审查的方案，不自动发布。",
            "clarification_questions": [],
            "deliverables": [{"name": "经营分析说明书", "content": detailed_content}],
            "implementation_steps": [],
            "artifacts": [],
            "test_checklist": [],
            "risks": [],
            "human_approval_points": [],
        }

        result = _validate_engineering_brief(payload, "power_bi")

        self.assertIn(detailed_content, result["deliverables"][0])

    def test_engineering_brief_allows_optional_sections_and_common_aliases(self) -> None:
        payload = {
            "schema_version": 1,
            "status": "ready",
            "category": "power_bi",
            "normalized_request": "设计 Power BI 工程方案",
            "outputs": ["星型模型", "核心 DAX 指标"],
            "steps": ["建立维表与事实表"],
            "acceptance_checklist": ["汇总值与源数据一致"],
            "approval_points": ["发布前人工审批"],
            "summary": "模型额外返回但本地无需执行的说明字段",
        }

        result = _validate_engineering_brief(payload, "power_bi")

        self.assertEqual(result["deliverables"], ["星型模型", "核心 DAX 指标"])
        self.assertEqual(result["implementation_steps"], ["建立维表与事实表"])
        self.assertEqual(result["artifacts"], [])
        self.assertEqual(result["risks"], [])


if __name__ == "__main__":
    unittest.main()
