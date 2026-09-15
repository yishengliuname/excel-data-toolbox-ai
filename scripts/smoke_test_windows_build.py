"""Start a packaged Windows app and verify the local demo/export path."""

from __future__ import annotations

import argparse
import json
import os
from pathlib import Path
import socket
import subprocess
import tempfile
import time
from urllib.error import URLError
from urllib.parse import urljoin
from urllib.request import Request, urlopen

from openpyxl import load_workbook


def free_port() -> int:
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        return int(sock.getsockname()[1])


def request_json(url: str, *, payload: dict | None = None, task_id: str = "") -> dict:
    data = None if payload is None else json.dumps(payload).encode("utf-8")
    headers = {"Accept": "application/json"}
    if data is not None:
        headers["Content-Type"] = "application/json"
    if task_id:
        headers["X-Task-ID"] = task_id
    request = Request(url, data=data, headers=headers, method="POST" if data is not None else "GET")
    with urlopen(request, timeout=20) as response:
        return json.loads(response.read().decode("utf-8"))


def wait_for_health(base_url: str, timeout: float = 45) -> dict:
    deadline = time.monotonic() + timeout
    last_error: Exception | None = None
    while time.monotonic() < deadline:
        try:
            return request_json(urljoin(base_url, "/health"))
        except (URLError, TimeoutError, ConnectionError) as exc:
            last_error = exc
            time.sleep(0.5)
    raise RuntimeError(f"packaged server did not become healthy: {last_error}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--app-dir", type=Path, required=True)
    parser.add_argument("--expected-version", required=True)
    args = parser.parse_args()

    app_dir = args.app_dir.resolve()
    executable = app_dir / "BiaogeKuaichuAI.exe"
    if not executable.is_file():
        raise SystemExit(f"missing packaged executable: {executable}")

    port = free_port()
    base_url = f"http://127.0.0.1:{port}"
    with tempfile.TemporaryDirectory(prefix="excel-toolbox-release-smoke-") as temp_name:
        temp_root = Path(temp_name)
        environment = os.environ.copy()
        environment["BIAOGE_USER_DATA"] = str(temp_root / "user-data")
        environment.pop("DEEPSEEK_API_KEY", None)
        process = subprocess.Popen(
            [str(executable), "--no-browser", "--port", str(port)],
            cwd=app_dir,
            env=environment,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
        try:
            health = wait_for_health(base_url)
            if health.get("status") != "ok" or health.get("version") != args.expected_version:
                raise RuntimeError(f"unexpected health response: {health}")

            with urlopen(urljoin(base_url, "/"), timeout=20) as response:
                page = response.read().decode("utf-8")
            if "表格快处" not in page or "加载演示数据" not in page:
                raise RuntimeError("packaged web UI did not load the expected application shell")

            demo = request_json(urljoin(base_url, "/api/demo"), payload={})
            task_id = str(demo["task_id"])
            table_id = str(demo["table_id"])
            state = request_json(urljoin(base_url, "/api/state"), task_id=task_id)
            if len(state.get("tables", [])) != 1 or state["tables"][0].get("rows", 0) <= 0:
                raise RuntimeError(f"synthetic demo did not load exactly one non-empty table: {state}")

            analysis = request_json(
                urljoin(base_url, "/api/analysis"), payload={"table": table_id}, task_id=task_id
            )
            if not analysis:
                raise RuntimeError("deterministic analysis returned an empty response")

            exported = request_json(
                urljoin(base_url, "/api/export"),
                payload={
                    "tables": [table_id],
                    "format": "xlsx",
                    "filename": "synthetic-release-smoke",
                    "include_summary": True,
                },
                task_id=task_id,
            )
            workbook_path = temp_root / "synthetic-release-smoke.xlsx"
            download_request = Request(
                urljoin(base_url, str(exported["download_url"])), headers={"X-Task-ID": task_id}
            )
            with urlopen(download_request, timeout=20) as response:
                workbook_path.write_bytes(response.read())
            workbook = load_workbook(workbook_path, read_only=True, data_only=False)
            try:
                if "处理摘要" not in workbook.sheetnames or len(workbook.sheetnames) < 2:
                    raise RuntimeError(f"exported workbook is incomplete: {workbook.sheetnames}")
            finally:
                workbook.close()
        finally:
            process.terminate()
            try:
                process.wait(timeout=10)
            except subprocess.TimeoutExpired:
                process.kill()
                process.wait(timeout=10)

    print("Windows package smoke test passed: health, UI, synthetic demo, analysis, export, reopen")


if __name__ == "__main__":
    main()
