"""Repeatable prompt -> validated plan -> values -> verified workbook checks.

Planner type is explicitly reported: mock tests never prove model comprehension.
Only IDs, failed check names and fingerprints are returned, not customer values.
"""
from __future__ import annotations

from dataclasses import dataclass
import math
from pathlib import Path
import re
from typing import Any, Callable, Mapping, Sequence

import pandas as pd
from openpyxl import load_workbook

from .core import export_tables
from .delivery_qa import dataframe_fingerprint, export_verified
from .nl_agent import AgentPlan, build_table_catalog, execute_plan, validate_plan


@dataclass(frozen=True)
class OutputCheck:
    table: str
    column: str
    expected: Any
    selector_column: str
    selector_value: Any
    absolute_tolerance: float = 1e-6

    def verify(self, tables: Mapping[str, pd.DataFrame]) -> bool:
        if self.absolute_tolerance < 0 or not math.isfinite(self.absolute_tolerance):
            raise ValueError("Invalid tolerance")
        table = tables.get(self.table)
        if table is None or self.column not in table or self.selector_column not in table:
            return False
        selected = table.loc[table[self.selector_column].eq(self.selector_value), self.column]
        if len(selected) != 1:
            return False
        value = selected.iloc[0]
        if isinstance(self.expected, (int, float)) and not isinstance(self.expected, bool):
            try:
                return math.isfinite(float(value)) and math.isclose(float(value), float(self.expected),
                    abs_tol=self.absolute_tolerance, rel_tol=0)
            except (TypeError, ValueError):
                return False
        return bool(pd.notna(value) and value == self.expected)


@dataclass(frozen=True)
class WorkflowScenario:
    scenario_id: str
    prompt: str
    expected_operations: tuple[str, ...]
    checks: tuple[OutputCheck, ...]
    required_tables: tuple[str, ...] = ()
    minimum_charts: int = 0


def run_workflow_evaluation(
    scenarios: Sequence[WorkflowScenario], tables: Mapping[str, pd.DataFrame], *,
    planner: Callable[[str, Mapping[str, Any]], Any],
    planner_kind: str, destination_dir: str | Path,
) -> dict[str, Any]:
    """Sequentially execute supplied acceptance cases; failed values never publish."""
    if planner_kind not in {"local", "mock", "remote"}:
        raise ValueError("Planner provenance must be local, mock or remote")
    if not scenarios or len({s.scenario_id for s in scenarios}) != len(scenarios):
        raise ValueError("Evaluation requires unique nonempty scenarios")
    if any(not re.fullmatch(r"[A-Za-z0-9_-]{3,80}", s.scenario_id) or not s.checks
           or not s.expected_operations or s.minimum_charts < 0 for s in scenarios):
        raise ValueError("Each scenario needs an ID, operations and output checks")
    before = {name: dataframe_fingerprint(frame) for name, frame in tables.items()}
    results = []
    destination = Path(destination_dir)
    for scenario in scenarios:
        failures = []
        workbook_verified = False
        artifacts = {}
        try:
            isolated = {name: frame.copy(deep=True) for name, frame in tables.items()}
            catalog = build_table_catalog(isolated)
            raw_plan = planner(scenario.prompt, catalog)
            payload = raw_plan.to_dict() if isinstance(raw_plan, AgentPlan) else raw_plan
            plan = validate_plan(payload, catalog)
            operations = {s.operation for s in plan.steps}
            failures.extend("missing_operation:" + op for op in scenario.expected_operations if op not in operations)
            if failures:
                raise ValueError("Plan does not match acceptance contract")
            execution = execute_plan(plan, isolated, dry_run=False)
            failures.extend("missing_table:" + name for name in scenario.required_tables if name not in execution.tables)
            failures.extend("output_check:" + str(i) for i, check in enumerate(scenario.checks, 1)
                            if not check.verify(execution.tables))
            if failures:
                raise ValueError("Business output checks failed")

            def verify_charts(path: Path) -> None:
                book = load_workbook(path, read_only=False, data_only=False)
                try:
                    if sum(len(sheet._charts) for sheet in book) < scenario.minimum_charts:
                        raise ValueError("Insufficient native charts")
                finally:
                    book.close()

            acceptance = export_verified(execution.tables, destination / f"{scenario.scenario_id}.xlsx",
                writer=lambda data, path: export_tables(data, path, include_log=False, overwrite=True), extra_validator=verify_charts)
            workbook_verified = acceptance.status == "passed"
            artifacts = {name: dataframe_fingerprint(frame) for name, frame in execution.tables.items()}
        except Exception as exc:
            # Do not echo an exception that may contain source values or paths.
            failures.append("execution_error:" + type(exc).__name__)
        if any(dataframe_fingerprint(frame) != before[name] for name, frame in tables.items()):
            failures.append("source_mutated")
        results.append({"scenario_id": scenario.scenario_id, "passed": not failures and workbook_verified,
                        "failures": failures, "workbook_verified": workbook_verified,
                        "output_fingerprints": artifacts})
    return {"schema_version": 1, "planner_kind": planner_kind,
            "scope": "plan, execution, business assertions, source immutability, workbook integrity, native charts",
            "passed": sum(r["passed"] for r in results), "failed": sum(not r["passed"] for r in results),
            "results": results}
