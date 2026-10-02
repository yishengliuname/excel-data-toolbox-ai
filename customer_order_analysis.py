"""Deterministic multi-table customer-order analysis.

The module converts a natural-language order-analysis request into a bounded
local workflow.  It does not execute model-generated code.  Sheet and column
roles are inferred from conservative aliases; ambiguous or missing evidence
stops execution instead of silently guessing.
"""

from __future__ import annotations

from collections import Counter
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
import math
import re
import unicodedata
from typing import Any

import pandas as pd
from openpyxl import load_workbook

from .status_semantics import ORDER_INVALID, ORDER_SUCCESS, UNKNOWN, classify_order_status


_HEADER_SEPARATORS = re.compile(r"[\s_\-()（）\[\]【】/:.：]+")
_TEXT_SPACES = re.compile(r"\s+")
_STORE_SEPARATORS = re.compile(r"[\s_\-—/，,。.]+")
_DIGIT_TO_CHINESE = str.maketrans(
    {"0": "零", "1": "一", "2": "二", "3": "三", "4": "四", "5": "五", "6": "六", "7": "七", "8": "八", "9": "九"}
)


_ALIASES: Mapping[str, tuple[str, ...]] = {
    "order_id": ("订单号", "订单编号", "交易单号", "orderid", "orderno", "order_number"),
    "order_date": ("下单时间", "下单日期", "交易时间", "交易日期", "销售日期", "date", "orderdate"),
    "store": ("门店", "门店名称", "店铺", "店铺名称", "门店编号", "store", "storename"),
    "product_code": ("商品编码", "商品代码", "sku", "sku编码", "货号", "productcode", "itemcode"),
    "product_name": ("商品名称", "标准商品名", "标准商品名称", "品名", "productname", "itemname"),
    "category": ("类目", "标准类目", "商品类别", "品类", "分类", "category"),
    "quantity": ("数量", "销售数量", "销量", "件数", "qty", "quantity"),
    "unit_price": ("单价", "销售单价", "成交单价", "price", "unitprice"),
    "discount": ("优惠金额", "折扣金额", "优惠", "discount", "discountamount"),
    "payment": ("支付方式", "付款方式", "paymentmethod", "payment"),
    "customer": ("会员手机号", "客户手机号", "客户编码", "会员编码", "customerid", "memberid"),
    "order_status": ("订单状态", "交易状态", "支付状态", "status", "orderstatus"),
    "cost_price": ("成本价", "单位成本", "标准成本", "采购成本", "cost", "unitcost"),
    "refund_amount": ("退款金额", "退款额", "退货金额", "refundamount", "refund"),
    "refund_date": ("退款时间", "退款日期", "refunddate"),
    "refund_reason": ("退款原因", "退货原因", "refundreason"),
    "target_amount": (
        "净销售额目标",
        "销售额目标",
        "销售目标",
        "营业额目标",
        "targetsales",
        "salestarget",
    ),
}


def _header_key(value: object) -> str:
    return _HEADER_SEPARATORS.sub("", unicodedata.normalize("NFKC", str(value))).casefold()


def _text(value: object) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    return _TEXT_SPACES.sub(" ", unicodedata.normalize("NFKC", str(value))).strip()


def _identifier(value: object) -> str:
    return _TEXT_SPACES.sub("", _text(value)).upper()


def _find_column(frame: pd.DataFrame, role: str) -> str | None:
    aliases = {_header_key(item) for item in _ALIASES[role]}
    matches = [str(column) for column in frame.columns if _header_key(column) in aliases]
    if len(matches) > 1:
        raise ValueError(f"字段角色 {role} 存在多个候选列：{'、'.join(matches)}")
    return matches[0] if matches else None


def _find_target_amount(frame: pd.DataFrame) -> str | None:
    direct = _find_column(frame, "target_amount")
    if direct:
        return direct
    candidates = [
        str(column)
        for column in frame.columns
        if "目标" in _header_key(column)
        and any(token in _header_key(column) for token in ("销售", "营业", "收入"))
        and "率" not in _header_key(column)
    ]
    return candidates[0] if len(candidates) == 1 else None


def _numeric(series: pd.Series) -> pd.Series:
    cleaned = series.map(
        lambda value: re.sub(r"[,，￥¥%\s]", "", _text(value)) if isinstance(value, str) else value
    )
    return pd.to_numeric(cleaned, errors="coerce")


def _parse_dates(series: pd.Series) -> pd.Series:
    parsed = pd.to_datetime(series, errors="coerce", format="mixed", dayfirst=False)
    if parsed.notna().mean() < 0.8:
        alternative = pd.to_datetime(series, errors="coerce", format="mixed", dayfirst=True)
        if alternative.notna().sum() > parsed.notna().sum():
            parsed = alternative
    return parsed


def _store_key(value: object) -> str:
    text = _STORE_SEPARATORS.sub("", _text(value)).casefold()
    text = text.replace("店铺", "店").replace("门店", "店")
    return text.translate(_DIGIT_TO_CHINESE)


def _unique_reference_map(values: pd.Series, key_function) -> tuple[dict[str, str], set[str]]:
    buckets: dict[str, set[str]] = {}
    for value in values:
        canonical = _text(value)
        if not canonical:
            continue
        buckets.setdefault(key_function(canonical), set()).add(canonical)
    ambiguous = {key for key, choices in buckets.items() if len(choices) != 1}
    return {key: next(iter(choices)) for key, choices in buckets.items() if len(choices) == 1}, ambiguous


@dataclass(frozen=True)
class OrderSourceBinding:
    sales_index: int
    master_index: int
    refund_index: int | None
    target_index: int | None
    columns: Mapping[str, str | None]


@dataclass(frozen=True)
class CustomerOrderAnalysisResult:
    outputs: Mapping[str, pd.DataFrame]
    report: Mapping[str, Any]


def infer_customer_order_sources(
    frames: Sequence[pd.DataFrame],
    *,
    source_names: Sequence[str] | None = None,
) -> OrderSourceBinding:
    """Infer order-detail, product-master, refund and target sources safely."""

    if not frames:
        raise ValueError("至少需要一张数据表")
    names = list(source_names or [f"表{index + 1}" for index in range(len(frames))])
    if len(names) != len(frames):
        raise ValueError("source_names 数量必须与表数量一致")

    role_columns: list[dict[str, str | None]] = []
    for frame in frames:
        role_columns.append({role: _find_column(frame, role) for role in _ALIASES})

    def choose(role: str, scorer) -> int | None:
        scores = [(scorer(columns), index) for index, columns in enumerate(role_columns)]
        best_score, best_index = max(scores, default=(0, -1))
        if best_score <= 0:
            return None
        tied = [index for score, index in scores if score == best_score]
        if len(tied) > 1:
            labels = "、".join(names[index] for index in tied)
            raise ValueError(f"无法唯一确定{role}数据源：{labels}")
        return best_index

    sales_required = ("order_id", "order_date", "store", "product_code", "quantity", "unit_price", "order_status")
    sales_index = choose(
        "销售明细",
        lambda c: sum(3 for role in sales_required if c[role])
        + sum(1 for role in ("discount", "payment", "customer", "product_name", "category") if c[role]),
    )
    if sales_index is None or any(not role_columns[sales_index][role] for role in sales_required):
        missing = [role for role in sales_required if sales_index is None or not role_columns[sales_index][role]]
        raise ValueError("无法安全执行订单分析，销售明细缺少关键字段：" + "、".join(missing))

    master_index = choose(
        "商品主数据",
        lambda c: (5 if c["product_code"] else 0)
        + (4 if c["cost_price"] else 0)
        + (2 if c["product_name"] else 0)
        + (2 if c["category"] else 0),
    )
    if master_index is None or master_index == sales_index:
        raise ValueError("未找到可用的独立商品主数据表")
    master_roles = role_columns[master_index]
    missing_master = [role for role in ("product_code", "product_name", "category", "cost_price") if not master_roles[role]]
    if missing_master:
        raise ValueError("商品主数据缺少必要字段：" + "、".join(missing_master))

    refund_candidates = [
        index
        for index, columns in enumerate(role_columns)
        if index not in {sales_index, master_index} and columns["order_id"] and columns["refund_amount"]
    ]
    if len(refund_candidates) > 1:
        raise ValueError("存在多张退款明细候选表，需要人工确认")
    refund_index = refund_candidates[0] if refund_candidates else None

    target_candidates = [
        index
        for index, columns in enumerate(role_columns)
        if index not in {sales_index, master_index, refund_index}
        and columns["store"]
        and _find_target_amount(frames[index])
    ]
    if len(target_candidates) > 1:
        raise ValueError("存在多张门店目标候选表，需要人工确认")
    target_index = target_candidates[0] if target_candidates else None

    sales_roles = role_columns[sales_index]
    columns: dict[str, str | None] = {f"sales_{key}": value for key, value in sales_roles.items()}
    columns.update({f"master_{key}": value for key, value in master_roles.items()})
    if refund_index is not None:
        columns.update({f"refund_{key}": value for key, value in role_columns[refund_index].items()})
    if target_index is not None:
        columns["target_store"] = role_columns[target_index]["store"]
        columns["target_amount"] = _find_target_amount(frames[target_index])
    return OrderSourceBinding(sales_index, master_index, refund_index, target_index, columns)


def can_build_customer_order_analysis(frames: Sequence[pd.DataFrame]) -> bool:
    try:
        infer_customer_order_sources(frames)
    except (TypeError, ValueError):
        return False
    return True


def _safe_master(frame: pd.DataFrame, binding: OrderSourceBinding) -> pd.DataFrame:
    code = str(binding.columns["master_product_code"])
    name = str(binding.columns["master_product_name"])
    category = str(binding.columns["master_category"])
    cost = str(binding.columns["master_cost_price"])
    master = pd.DataFrame(
        {
            "商品编码": frame[code].map(_identifier),
            "主数据商品名称": frame[name].map(_text),
            "主数据类目": frame[category].map(_text),
            "主数据成本价": _numeric(frame[cost]),
        }
    )
    master = master.loc[master["商品编码"].ne("")].copy()
    conflicts: list[str] = []
    for code_value, group in master.groupby("商品编码", dropna=False):
        signatures = group[["主数据商品名称", "主数据类目", "主数据成本价"]].drop_duplicates()
        if len(signatures) > 1:
            conflicts.append(str(code_value))
    if conflicts:
        raise ValueError("商品主数据关键字段冲突，无法执行 N:1 关联：" + "、".join(conflicts[:20]))
    master = master.drop_duplicates(subset=["商品编码"], keep="first")
    if master["主数据成本价"].isna().any() or master["主数据成本价"].lt(0).any():
        bad = master.loc[
            master["主数据成本价"].isna() | master["主数据成本价"].lt(0),
            "商品编码",
        ].tolist()
        raise ValueError("商品主数据成本价缺失或小于0：" + "、".join(map(str, bad[:20])))
    return master.reset_index(drop=True)


def _status_reason(value: object, classification: str) -> str:
    text = _text(value)
    compact = _HEADER_SEPARATORS.sub("", text)
    if "取消" in compact or "作废" in compact or "关闭" in compact:
        return f"订单状态={text or '无法识别'}"
    if any(token in compact for token in ("待支付", "待付款", "未支付")):
        return f"订单状态={text or '无法识别'}"
    if classification == ORDER_INVALID:
        return f"订单状态={text or '无法识别'}"
    return "订单状态无法确认"


def _append_excluded(
    rows: list[pd.DataFrame],
    source: pd.DataFrame,
    mask: pd.Series,
    reason: str | pd.Series,
) -> None:
    if not bool(mask.any()):
        return
    part = source.loc[mask].copy()
    part.insert(0, "异常类型", "业务排除")
    if isinstance(reason, pd.Series):
        part.insert(1, "排除原因", reason.loc[mask].astype("string").values)
    else:
        part.insert(1, "排除原因", reason)
    rows.append(part)


def _audit_row(
    operation: str,
    before: int,
    after: int,
    affected: int,
    detail: str,
    *,
    status: str = "通过",
) -> dict[str, Any]:
    return {
        "顺序": 0,
        "执行操作": operation,
        "执行前行数": before,
        "执行后行数": after,
        "影响行数": affected,
        "结果": status,
        "说明": detail,
    }


def build_customer_order_analysis(
    frames: Sequence[pd.DataFrame],
    *,
    source_names: Sequence[str] | None = None,
    user_request: str = "",
    low_margin_threshold: float = 0.20,
) -> CustomerOrderAnalysisResult:
    """Execute a validated order-analysis workflow and return report tables."""

    if not 0 <= float(low_margin_threshold) <= 1:
        raise ValueError("low_margin_threshold 必须在0到1之间")
    local_frames = [frame.copy(deep=True) for frame in frames]
    names = list(source_names or [f"表{index + 1}" for index in range(len(local_frames))])
    binding = infer_customer_order_sources(local_frames, source_names=names)
    sales_source = local_frames[binding.sales_index].copy(deep=True)
    master = _safe_master(local_frames[binding.master_index], binding)
    audit: list[dict[str, Any]] = []
    warnings: list[str] = []

    sales_columns = {
        key.replace("sales_", ""): value
        for key, value in binding.columns.items()
        if key.startswith("sales_") and value
    }
    raw_rows = len(sales_source)
    original = sales_source.copy(deep=True)
    original.insert(0, "源行号", range(2, len(original) + 2))
    working = original.copy(deep=True)
    for column in sales_source.columns:
        working[column] = working[column].map(_text)
    code_column = str(sales_columns["product_code"])
    order_column = str(sales_columns["order_id"])
    working[code_column] = working[code_column].map(_identifier)
    working[order_column] = working[order_column].map(_identifier)
    audit.append(_audit_row("normalize_text", raw_rows, raw_rows, raw_rows, "NFKC、首尾空格和连续空白标准化"))
    audit.append(_audit_row("case_normalize", raw_rows, raw_rows, raw_rows, "订单号和商品编码统一为大写标识符"))

    compare_columns = list(sales_source.columns)
    duplicate_mask = working[compare_columns].duplicated(keep="first")
    duplicate_rows = int(duplicate_mask.sum())
    excluded_parts: list[pd.DataFrame] = []
    _append_excluded(excluded_parts, original, duplicate_mask, "标准化后完全重复")
    working = working.loc[~duplicate_mask].copy()
    original = original.loc[~duplicate_mask].copy()
    deduplicated_rows = len(working)
    audit.append(_audit_row("deduplicate", raw_rows, deduplicated_rows, duplicate_rows, "以标准化后的全部原始字段识别完全重复"))

    target_frame: pd.DataFrame | None = None
    target_map: dict[str, float] = {}
    target_store_map: dict[str, str] = {}
    ambiguous_store_keys: set[str] = set()
    if binding.target_index is not None:
        target_frame = local_frames[binding.target_index].copy(deep=True)
        target_store_column = str(binding.columns["target_store"])
        target_amount_column = str(binding.columns["target_amount"])
        target_store_map, ambiguous_store_keys = _unique_reference_map(target_frame[target_store_column], _store_key)
        target_values = _numeric(target_frame[target_amount_column])
        for store_value, amount in zip(target_frame[target_store_column], target_values, strict=True):
            key = _store_key(store_value)
            if key in target_store_map and pd.notna(amount):
                canonical = target_store_map[key]
                if canonical in target_map and not math.isclose(target_map[canonical], float(amount)):
                    raise ValueError(f"门店目标存在冲突：{canonical}")
                target_map[canonical] = float(amount)
    else:
        warnings.append("未找到门店目标表；目标和完成率将标记为不可用")

    store_column = str(sales_columns["store"])
    if target_store_map:
        store_keys = working[store_column].map(_store_key)
        working["门店_标准"] = store_keys.map(target_store_map)
        working.loc[store_keys.isin(ambiguous_store_keys), "门店_标准"] = pd.NA
    else:
        working["门店_标准"] = working[store_column].map(_text).replace("", pd.NA)
    changed_store = int(
        (
            working["门店_标准"].notna()
            & working["门店_标准"].astype("string").ne(working[store_column].astype("string"))
        ).sum()
    )
    audit.append(
        _audit_row(
            "reference_mapping",
            len(working),
            len(working),
            changed_store,
            "门店别名仅在与门店目标参考表形成唯一安全匹配时自动转换",
        )
    )

    master_codes = set(master["商品编码"])
    status_column = str(sales_columns["order_status"])
    quantity_column = str(sales_columns["quantity"])
    price_column = str(sales_columns["unit_price"])
    date_column = str(sales_columns["order_date"])
    status_class = working[status_column].map(classify_order_status)
    quantity = _numeric(working[quantity_column])
    unit_price = _numeric(working[price_column])
    parsed_date = _parse_dates(working[date_column])

    reason = pd.Series("", index=working.index, dtype="object")
    invalid_status = status_class.ne(ORDER_SUCCESS)
    reason.loc[invalid_status] = [
        _status_reason(value, classification)
        for value, classification in zip(
            working.loc[invalid_status, status_column], status_class.loc[invalid_status], strict=True
        )
    ]
    invalid_store = reason.eq("") & working["门店_标准"].isna()
    reason.loc[invalid_store] = "门店无法识别"
    invalid_product = reason.eq("") & ~working[code_column].isin(master_codes)
    reason.loc[invalid_product] = "商品编码无法匹配主数据"
    invalid_quantity = reason.eq("") & (quantity.isna() | quantity.le(0))
    reason.loc[invalid_quantity] = "数量<=0"
    invalid_price = reason.eq("") & (unit_price.isna() | unit_price.le(0))
    reason.loc[invalid_price] = "单价<=0"
    invalid_date = reason.eq("") & parsed_date.isna()
    reason.loc[invalid_date] = "下单时间无法解析"

    business_invalid = reason.ne("")
    original_for_reason = original.loc[working.index]
    _append_excluded(excluded_parts, original_for_reason, business_invalid, reason)
    reason_counts = Counter(reason.loc[business_invalid].tolist())
    audit.append(
        _audit_row(
            "exclude_with_reason",
            len(working),
            int((~business_invalid).sum()),
            int(business_invalid.sum()),
            "；".join(f"{key} {value}条" for key, value in reason_counts.items()),
        )
    )

    valid = working.loc[~business_invalid].copy()
    valid["数量_数值"] = quantity.loc[valid.index]
    valid["单价_数值"] = unit_price.loc[valid.index]
    valid["下单时间_标准"] = parsed_date.loc[valid.index]
    before_join = len(valid)
    valid = valid.merge(master, left_on=code_column, right_on="商品编码", how="left", validate="many_to_one")
    if len(valid) != before_join:
        raise ValueError("商品主数据关联导致行数变化，已停止交付")
    if valid[["主数据商品名称", "主数据类目", "主数据成本价"]].isna().any().any():
        raise ValueError("商品主数据关联不完整，已停止交付")
    audit.append(_audit_row("safe_join", before_join, len(valid), len(valid), "销售明细 N:1 关联商品主数据；关联前后行数一致"))

    discount_column = sales_columns.get("discount")
    discount = _numeric(valid[str(discount_column)]) if discount_column else pd.Series(0.0, index=valid.index)
    discount = discount.fillna(0.0)
    invalid_discount = discount.lt(0)
    if bool(invalid_discount.any()):
        raise ValueError("优惠金额存在负数，无法在未确认业务口径时继续")
    valid["销售额"] = (valid["数量_数值"] * valid["单价_数值"]).round(2)
    valid["优惠金额_数值"] = discount.round(2)

    refund_total_by_order = pd.Series(dtype="float64")
    refund_review_parts: list[pd.DataFrame] = []
    refund_source_rows = 0
    if binding.refund_index is not None:
        refund_source = local_frames[binding.refund_index].copy(deep=True)
        refund_source_rows = len(refund_source)
        refund_order_column = str(binding.columns["refund_order_id"])
        refund_amount_column = str(binding.columns["refund_refund_amount"])
        refund_work = refund_source.copy(deep=True)
        refund_work["订单号_标准"] = refund_work[refund_order_column].map(_identifier)
        refund_work["退款金额_数值"] = _numeric(refund_work[refund_amount_column])
        invalid_refund = refund_work["订单号_标准"].eq("") | refund_work["退款金额_数值"].isna() | refund_work["退款金额_数值"].lt(0)
        if bool(invalid_refund.any()):
            review = refund_work.loc[invalid_refund].copy()
            review.insert(0, "人工核验原因", "退款订单号或退款金额无效")
            refund_review_parts.append(review)
        usable_refunds = refund_work.loc[~invalid_refund].copy()
        refund_total_by_order = usable_refunds.groupby("订单号_标准", dropna=False)["退款金额_数值"].sum()
        valid_orders_set = set(valid[order_column])
        unmatched_refund = usable_refunds[~usable_refunds["订单号_标准"].isin(valid_orders_set)]
        if not unmatched_refund.empty:
            review = unmatched_refund.copy()
            review.insert(0, "人工核验原因", "退款订单未进入有效经营明细（可能为已排除订单）")
            refund_review_parts.append(review)
    elif "退款" in user_request:
        raise ValueError("用户要求考虑退款，但未找到包含订单号和退款金额的退款表")

    order_gross = valid.groupby(order_column)["销售额"].transform("sum")
    order_refund = valid[order_column].map(refund_total_by_order).fillna(0.0)
    valid["退款金额"] = (order_refund * valid["销售额"] / order_gross).round(2)
    valid["净销售额"] = (valid["销售额"] - valid["优惠金额_数值"] - valid["退款金额"]).round(2)
    valid["成本"] = (valid["数量_数值"] * valid["主数据成本价"]).round(2)
    valid["毛利"] = (valid["净销售额"] - valid["成本"]).round(2)
    valid["毛利率"] = valid["毛利"].div(valid["净销售额"].where(valid["净销售额"].ne(0)))
    audit.append(
        _audit_row(
            "aggregate_refund",
            refund_source_rows,
            len(refund_total_by_order),
            refund_source_rows,
            "同一订单的多条退款先汇总，再按有效商品销售额比例分摊并逐行保留两位小数",
        )
    )
    audit.append(_audit_row("derive_column", len(valid), len(valid), len(valid), "计算销售额、净销售额、成本、毛利和毛利率"))

    cleaned = pd.DataFrame(
        {
            "源行号": valid["源行号"],
            "订单号": valid[order_column],
            "下单时间": valid["下单时间_标准"],
            "日期": valid["下单时间_标准"].dt.normalize(),
            "门店": valid["门店_标准"],
            "商品编码": valid[code_column],
            "商品名称": valid["主数据商品名称"],
            "类目": valid["主数据类目"],
            "数量": valid["数量_数值"],
            "单价": valid["单价_数值"],
            "优惠金额": valid["优惠金额_数值"],
            "支付方式": valid[str(sales_columns["payment"])] if sales_columns.get("payment") else "",
            "客户标识": valid[str(sales_columns["customer"])] if sales_columns.get("customer") else "",
            "订单状态": valid[status_column],
            "销售额": valid["销售额"],
            "退款金额": valid["退款金额"],
            "净销售额": valid["净销售额"],
            "成本价": valid["主数据成本价"],
            "成本": valid["成本"],
            "毛利": valid["毛利"],
            "毛利率": valid["毛利率"],
        }
    ).reset_index(drop=True)

    store_group = cleaned.groupby("门店", dropna=False)
    store_summary = store_group.agg(
        销售额=("销售额", "sum"),
        优惠金额=("优惠金额", "sum"),
        退款金额=("退款金额", "sum"),
        净销售额=("净销售额", "sum"),
        成本=("成本", "sum"),
        毛利=("毛利", "sum"),
        有效订单数=("订单号", "nunique"),
    ).reset_index()
    store_summary["目标"] = store_summary["门店"].map(target_map)
    store_summary["目标完成率"] = store_summary["净销售额"].div(store_summary["目标"].where(store_summary["目标"].ne(0)))
    store_summary["毛利率"] = store_summary["毛利"].div(store_summary["净销售额"].where(store_summary["净销售额"].ne(0)))
    store_summary["客单价"] = store_summary["净销售额"].div(store_summary["有效订单数"].where(store_summary["有效订单数"].ne(0)))
    store_summary["退款率"] = store_summary["退款金额"].div(store_summary["销售额"].where(store_summary["销售额"].ne(0)))
    store_summary = store_summary[
        ["门店", "销售额", "优惠金额", "退款金额", "净销售额", "目标", "目标完成率", "成本", "毛利", "毛利率", "有效订单数", "客单价", "退款率"]
    ].sort_values("净销售额", ascending=False, kind="stable").reset_index(drop=True)

    product_summary = cleaned.groupby(["商品编码", "商品名称", "类目"], dropna=False).agg(
        销量=("数量", "sum"),
        销售额=("销售额", "sum"),
        优惠金额=("优惠金额", "sum"),
        退款金额=("退款金额", "sum"),
        净销售额=("净销售额", "sum"),
        成本=("成本", "sum"),
        毛利=("毛利", "sum"),
        有效订单数=("订单号", "nunique"),
    ).reset_index()
    product_summary["毛利率"] = product_summary["毛利"].div(product_summary["净销售额"].where(product_summary["净销售额"].ne(0)))
    product_summary["低毛利标记"] = product_summary["毛利率"].lt(float(low_margin_threshold)).map({True: "需关注", False: ""})
    product_summary = product_summary.sort_values("净销售额", ascending=False, kind="stable").reset_index(drop=True)

    daily = cleaned.groupby("日期", dropna=False).agg(
        销售额=("销售额", "sum"),
        优惠金额=("优惠金额", "sum"),
        退款金额=("退款金额", "sum"),
        净销售额=("净销售额", "sum"),
        成本=("成本", "sum"),
        毛利=("毛利", "sum"),
        有效订单数=("订单号", "nunique"),
    ).reset_index().sort_values("日期").reset_index(drop=True)
    daily["毛利率"] = daily["毛利"].div(daily["净销售额"].where(daily["净销售额"].ne(0)))

    warning_parts: list[pd.DataFrame] = []
    for column in ("数量", "单价"):
        series = pd.to_numeric(cleaned[column], errors="coerce")
        if series.notna().sum() < 4:
            continue
        q1, q3 = series.quantile([0.25, 0.75])
        iqr = q3 - q1
        if iqr <= 0:
            continue
        lower, upper = q1 - 1.5 * iqr, q3 + 1.5 * iqr
        mask = series.lt(lower) | series.gt(upper)
        if bool(mask.any()):
            part = cleaned.loc[mask, ["源行号", "订单号", "门店", "商品编码", column]].copy()
            part.insert(0, "预警类型", "统计预警")
            part.insert(1, "预警原因", f"{column}超出IQR参考区间[{lower:.2f}, {upper:.2f}]；未从经营数据删除")
            warning_parts.append(part)
    statistical_warnings = pd.concat(warning_parts, ignore_index=True) if warning_parts else pd.DataFrame(
        columns=["预警类型", "预警原因", "源行号", "订单号", "门店", "商品编码", "数值"]
    )
    audit.append(_audit_row("statistical_warning", len(cleaned), len(cleaned), len(statistical_warnings), "IQR只生成参考预警，不自动排除有效记录"))

    excluded = pd.concat(excluded_parts, ignore_index=True, sort=False) if excluded_parts else pd.DataFrame()
    manual_review_parts: list[pd.DataFrame] = []
    review_reason_mask = excluded.get("排除原因", pd.Series(dtype="string")).isin(
        ["门店无法识别", "商品编码无法匹配主数据", "订单状态无法确认", "下单时间无法解析"]
    )
    if not excluded.empty and bool(review_reason_mask.any()):
        review = excluded.loc[review_reason_mask].copy()
        review.insert(0, "人工核验原因", review["排除原因"])
        manual_review_parts.append(review)
    manual_review_parts.extend(refund_review_parts)
    manual_review = pd.concat(manual_review_parts, ignore_index=True, sort=False) if manual_review_parts else pd.DataFrame(
        columns=["人工核验原因", "数据源"]
    )

    net_sales = round(float(cleaned["净销售额"].sum()), 2)
    gross_profit = round(float(cleaned["毛利"].sum()), 2)
    refund_amount = round(float(cleaned["退款金额"].sum()), 2)
    gross_margin = gross_profit / net_sales if net_sales else math.nan
    valid_orders = int(cleaned["订单号"].nunique())
    excluded_rows = len(excluded)
    valid_rows = len(cleaned)

    checks = [
        ("原始行数勾稽", raw_rows == deduplicated_rows + duplicate_rows, f"{raw_rows}={deduplicated_rows}+{duplicate_rows}"),
        ("去重后分流勾稽", deduplicated_rows == valid_rows + (excluded_rows - duplicate_rows), f"{deduplicated_rows}={valid_rows}+{excluded_rows - duplicate_rows}"),
        ("商品主数据匹配", cleaned["商品名称"].notna().all(), "有效明细商品名称非空"),
        ("门店标准化", cleaned["门店"].notna().all(), "有效明细门店非空"),
        ("明细汇总勾稽", math.isclose(net_sales, round(float(store_summary["净销售额"].sum()), 2), abs_tol=0.01), "门店汇总与明细一致"),
        ("关键KPI有效", all(math.isfinite(value) for value in (net_sales, gross_profit, gross_margin, refund_amount)), "净销售额、毛利、毛利率和退款金额均可计算"),
    ]
    failed_checks = [name for name, passed, _ in checks if not bool(passed)]
    if failed_checks:
        raise ValueError("数据验收未通过：" + "、".join(failed_checks))
    validation = pd.DataFrame(
        [{"验收项": name, "结果": "PASS" if passed else "FAIL", "证据": evidence} for name, passed, evidence in checks]
    )
    audit.append(_audit_row("validate", valid_rows, valid_rows, len(checks), "行数、关联、类别和金额勾稽均通过"))
    for index, row in enumerate(audit, start=1):
        row["顺序"] = index
    audit_frame = pd.DataFrame(audit)

    dashboard = pd.DataFrame(
        [
            {"指标": "净销售额", "结果": net_sales, "单位": "元", "数据口径": "销售额-优惠金额-有效退款分摊"},
            {"指标": "毛利", "结果": gross_profit, "单位": "元", "数据口径": "净销售额-商品主数据成本"},
            {"指标": "毛利率", "结果": gross_margin, "单位": "%", "数据口径": "毛利/净销售额"},
            {"指标": "有效订单数", "结果": valid_orders, "单位": "单", "数据口径": "有效经营明细的不同订单号"},
            {"指标": "有效退款金额", "结果": refund_amount, "单位": "元", "数据口径": "仅分摊到有效经营订单的退款"},
            {"指标": "有效明细行数", "结果": valid_rows, "单位": "行", "数据口径": "去重并执行业务排除后"},
            {"指标": "异常/排除记录", "结果": excluded_rows, "单位": "行", "数据口径": "含标准化后重复、无效状态和明确业务错误"},
        ]
    )

    risk_counts = excluded["排除原因"].value_counts().rename_axis("异常类型").reset_index(name="风险数量")
    if not statistical_warnings.empty:
        warning_labels = statistical_warnings["预警原因"].astype("string").str.extract(r"^([^\[]+)", expand=False)
        warning_counts = (
            ("统计预警：" + warning_labels.str.replace("超出IQR参考区间", "", regex=False).str.strip())
            .value_counts()
            .rename_axis("异常类型")
            .reset_index(name="风险数量")
        )
        risk_counts = pd.concat([risk_counts, warning_counts], ignore_index=True)
    chart_rows = max(len(daily), len(store_summary), len(product_summary.head(10)), len(risk_counts))
    chart_data = pd.DataFrame(index=range(chart_rows))
    chart_data["日期"] = daily["日期"].reindex(chart_data.index)
    chart_data["每日净销售额"] = daily["净销售额"].reindex(chart_data.index)
    chart_data["每日毛利"] = daily["毛利"].reindex(chart_data.index)
    chart_data["门店"] = store_summary["门店"].reindex(chart_data.index)
    chart_data["门店净销售额"] = store_summary["净销售额"].reindex(chart_data.index)
    top_products = product_summary.head(10).reset_index(drop=True)
    chart_data["商品"] = top_products["商品名称"].reindex(chart_data.index)
    chart_data["商品净销售额"] = top_products["净销售额"].reindex(chart_data.index)
    chart_data["异常类型"] = risk_counts["异常类型"].reindex(chart_data.index)
    chart_data["风险数量"] = risk_counts["风险数量"].reindex(chart_data.index)

    outputs = {
        "管理看板": dashboard,
        "门店汇总": store_summary,
        "商品分析": product_summary,
        "每日趋势": daily,
        "清洗明细": cleaned,
        "异常数据": excluded,
        "统计预警": statistical_warnings,
        "人工核验": manual_review,
        "执行审计": audit_frame,
        "数据验收": validation,
        "图表看板": chart_data,
    }
    for frame in outputs.values():
        frame.attrs["toolbox_report_kind"] = "customer_order_analysis"
    report = {
        "raw_rows": raw_rows,
        "duplicate_rows": duplicate_rows,
        "deduplicated_rows": deduplicated_rows,
        "valid_rows": valid_rows,
        "excluded_rows": excluded_rows,
        "net_sales": net_sales,
        "gross_profit": gross_profit,
        "gross_margin": gross_margin,
        "valid_orders": valid_orders,
        "refund_amount": refund_amount,
        "stores": sorted(cleaned["门店"].dropna().astype(str).unique().tolist()),
        "source_roles": {
            "sales": names[binding.sales_index],
            "product_master": names[binding.master_index],
            "refund": names[binding.refund_index] if binding.refund_index is not None else "",
            "target": names[binding.target_index] if binding.target_index is not None else "",
        },
        "columns": dict(binding.columns),
        "exclusion_reasons": dict(Counter(excluded["排除原因"].tolist())),
        "warnings": warnings,
        "operation_count": len(audit_frame),
        "validation_status": "passed",
    }
    return CustomerOrderAnalysisResult(outputs, report)


def validate_customer_order_params(params: Mapping[str, Any]) -> None:
    allowed = {"source_names", "user_request", "low_margin_threshold"}
    unknown = sorted(set(params) - allowed)
    if unknown:
        raise ValueError("订单分析参数不支持：" + "、".join(unknown))
    source_names = params.get("source_names")
    if not isinstance(source_names, (list, tuple)) or not source_names or not all(isinstance(item, str) and item for item in source_names):
        raise ValueError("source_names 必须是非空字符串列表")
    if "user_request" in params and not isinstance(params["user_request"], str):
        raise ValueError("user_request 必须是字符串")
    if "low_margin_threshold" in params:
        value = params["low_margin_threshold"]
        if isinstance(value, bool) or not isinstance(value, (int, float)) or not 0 <= float(value) <= 1:
            raise ValueError("low_margin_threshold 必须在0到1之间")


def verify_customer_order_workbook(path: str) -> Mapping[str, Any]:
    """Block delivery when the native workbook contract is incomplete."""

    workbook = load_workbook(path, data_only=False, read_only=False)
    required = tuple(_ORDER_DELIVERY_SHEETS)
    missing = [name for name in required if name not in workbook.sheetnames]
    if missing:
        raise ValueError("订单经营报告缺少工作表：" + "、".join(missing))
    chart_sheet = workbook["图表看板"]
    if len(chart_sheet._charts) < 4:
        raise ValueError("订单经营报告图表不完整，至少需要4个原生Excel图表")
    for chart in chart_sheet._charts:
        if not chart.series:
            raise ValueError("订单经营报告存在无数据系列的空图表")
    formula_errors = {"#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#N/A", "#NUM!", "#NULL!"}
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and any(token in cell.value for token in formula_errors):
                    raise ValueError(f"工作表 {worksheet.title} 含公式错误：{cell.coordinate}")
    cleaned = pd.read_excel(path, sheet_name="清洗明细", header=3)
    stores = pd.read_excel(path, sheet_name="门店汇总", header=3)
    excluded = pd.read_excel(path, sheet_name="异常数据", header=3)
    if cleaned.empty or stores.empty:
        raise ValueError("订单经营报告的有效明细或门店汇总为空")
    if cleaned[["门店", "商品编码", "商品名称", "类目"]].isna().any().any():
        raise ValueError("有效明细仍存在未标准化的关键业务字段")
    detail_net = round(float(pd.to_numeric(cleaned["净销售额"], errors="coerce").sum()), 2)
    store_net = round(float(pd.to_numeric(stores["净销售额"], errors="coerce").sum()), 2)
    if not math.isclose(detail_net, store_net, abs_tol=0.01):
        raise ValueError("有效明细与门店汇总净销售额不一致")
    if not excluded.empty and excluded["排除原因"].isna().any():
        raise ValueError("异常数据存在无排除原因记录")
    return {
        "status": "passed",
        "sheet_count": len(workbook.sheetnames),
        "chart_count": len(chart_sheet._charts),
        "cleaned_rows": len(cleaned),
        "excluded_rows": len(excluded),
        "net_sales": detail_net,
    }


_ORDER_DELIVERY_SHEETS = (
    "管理看板",
    "门店汇总",
    "商品分析",
    "每日趋势",
    "清洗明细",
    "异常数据",
    "统计预警",
    "人工核验",
    "执行审计",
    "数据验收",
    "图表看板",
)


__all__ = [
    "CustomerOrderAnalysisResult",
    "OrderSourceBinding",
    "build_customer_order_analysis",
    "can_build_customer_order_analysis",
    "infer_customer_order_sources",
    "validate_customer_order_params",
    "verify_customer_order_workbook",
]
